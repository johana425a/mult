"""DeepSteam bounded local document reader (our integration, not upstream code).

One JSON request on stdin, one JSON response on stdout. The caller authorizes
``root`` before spawning this helper; this module independently confines the
relative ``path`` to that root. No shell commands, network, original-file writes,
user-content logging, or disk cache. Run in the packaged Python with ``-I``.

read: {action, root, path, cursor?, max_chars?, max_units?, ocr?: auto|off}
asset: {action, root, path, asset_id, document_sha256, max_dimension?}
Pagination cursors and image reads are bound to the original document SHA256.
PDF page numbers are physical PDF pages, never asserted printed-book pages.
Office positions are paragraphs/tables/slides, not fictitious paper pages.
"""

from __future__ import annotations

import base64
import contextlib
import hashlib
import io
import json
import logging
import os
from pathlib import Path, PurePosixPath, PureWindowsPath
import posixpath
import re
import stat
import sys
from typing import Any
import warnings
import zipfile
from urllib.parse import unquote, urlsplit
from unittest import mock

SCHEMA_VERSION = 1
try:
    MAX_FILE_BYTES = min(512 * 1024 * 1024, max(1024 * 1024, int(os.environ.get("EDUAGENT_DOCUMENT_MAX_BYTES", str(200 * 1024 * 1024)))))
except ValueError:
    MAX_FILE_BYTES = 200 * 1024 * 1024
MAX_TEXT_BYTES = 32 * 1024 * 1024
MAX_ZIP_MEMBERS = 10000
MAX_ZIP_EXPANDED = 256 * 1024 * 1024
MAX_ZIP_MEMBER = 32 * 1024 * 1024
MAX_ZIP_RATIO = 200
MAX_UNITS = 100000
MAX_UNIT_CHARS = 2_000_000
MAX_ASSET_BYTES = 8 * 1024 * 1024
MAX_IMAGE_PIXELS = 16_000_000
MAX_IMAGES_PER_UNIT = 32
TEXT_FORMATS = {".md", ".markdown", ".txt", ".csv", ".tsv", ".json", ".rst", ".log"}
OFFICE_FORMATS = {".docx": "word/document.xml", ".pptx": "ppt/presentation.xml", ".xlsx": "xl/workbook.xml"}
IMAGE_FORMATS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
SUPPORTED_FORMATS = TEXT_FORMATS | set(OFFICE_FORMATS) | IMAGE_FORMATS | {".pdf"}
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MAX_EQUATION_CHARS = 16000


def omml_equation(element: Any) -> dict[str, Any]:
    """Readable structural conversion, not a Word layout renderer.

    Unsupported OMML constructs are retained as linear text and marked partial.
    Never execute or insert input TeX commands: m:t is escaped as literal math.
    """
    unresolved: set[str] = set()
    symbols = {"α": r"\alpha ", "β": r"\beta ", "γ": r"\gamma ", "δ": r"\delta ",
               "θ": r"\theta ", "λ": r"\lambda ", "μ": r"\mu ", "π": r"\pi ",
               "σ": r"\sigma ", "φ": r"\phi ", "ω": r"\omega ", "Δ": r"\Delta ",
               "∞": r"\infty ", "×": r"\times ", "÷": r"\div ", "≤": r"\leq ",
               "≥": r"\geq ", "≠": r"\neq ", "→": r"\to ", "−": "-"}

    def literal(value: str) -> str:
        special = {"\\": r"\backslash ", "{": r"\{", "}": r"\}", "#": r"\#", "%": r"\%", "&": r"\&", "_": r"\_", "$": r"\$", "^": r"\hat{}", "~": r"\sim "}
        return "".join(symbols.get(char, special.get(char, char)) for char in value)

    def prop(node: Any, group: str, name: str, default: str = "") -> str:
        item = node.find(f"{{{MATH_NS}}}{group}/{{{MATH_NS}}}{name}")
        return item.get(f"{{{MATH_NS}}}val", default) if item is not None else default

    def render(node: Any, depth: int = 0) -> str:
        if depth > 64:
            unresolved.add("equation_depth_limit")
            return r"\text{[nested equation]}"
        name = str(node.tag).rsplit("}", 1)[-1]
        if name.endswith("Pr"):
            return ""
        if name == "t":
            return literal(node.text or "")
        def child(tag: str) -> str:
            value = node.find(f"{{{MATH_NS}}}{tag}")
            return render(value, depth + 1) if value is not None else ""
        def joined() -> str:
            return "".join(render(item, depth + 1) for item in node)
        if name == "f":
            if prop(node, "fPr", "type") == "noBar":
                return r"\genfrac{}{}{0pt}{}{" + child("num") + "}{" + child("den") + "}"
            return r"\frac{" + child("num") + "}{" + child("den") + "}"
        if name in {"sSup", "sSub", "sSubSup"}:
            value = "{" + child("e") + "}"
            if name in {"sSub", "sSubSup"}:
                value += "_{" + child("sub") + "}"
            if name in {"sSup", "sSubSup"}:
                value += "^{" + child("sup") + "}"
            return value
        if name == "rad":
            degree = child("deg")
            return r"\sqrt" + ("[" + degree + "]" if degree else "") + "{" + child("e") + "}"
        if name == "nary":
            operators = {"∑": r"\sum", "∏": r"\prod", "∫": r"\int", "∬": r"\iint", "∭": r"\iiint", "∮": r"\oint", "⋃": r"\bigcup", "⋂": r"\bigcap"}
            char = prop(node, "naryPr", "chr", "∫")
            value = operators.get(char, literal(char))
            if child("sub"):
                value += "_{" + child("sub") + "}"
            if child("sup"):
                value += "^{" + child("sup") + "}"
            return value + " {" + child("e") + "}"
        if name == "d":
            delimiters = {"{": r"\{", "}": r"\}", "": ".", "⌊": r"\lfloor", "⌋": r"\rfloor", "⌈": r"\lceil", "⌉": r"\rceil", "⟨": r"\langle", "⟩": r"\rangle", "|": "|", "‖": r"\|", "(": "(", ")": ")", "[": "[", "]": "]"}
            begin, end = prop(node, "dPr", "begChr", "("), prop(node, "dPr", "endChr", ")")
            if begin not in delimiters or end not in delimiters:
                unresolved.add("unhandled_math_delimiter")
                return literal(begin) + joined() + literal(end)
            separator = literal(prop(node, "dPr", "sepChr", "|"))
            values = [render(item, depth + 1) for item in node if str(item.tag) == f"{{{MATH_NS}}}e"]
            return r"\left" + delimiters[begin] + separator.join(values) + r"\right" + delimiters[end]
        if name == "m":
            rows = [render(item, depth + 1) for item in node if str(item.tag) == f"{{{MATH_NS}}}mr"]
            return r"\begin{matrix}" + r" \\ ".join(rows) + r"\end{matrix}"
        if name == "mr":
            return " & ".join(render(item, depth + 1) for item in node if str(item.tag) == f"{{{MATH_NS}}}e")
        if name == "eqArr":
            return r"\begin{gathered}" + r" \\ ".join(render(item, depth + 1) for item in node if str(item.tag) == f"{{{MATH_NS}}}e") + r"\end{gathered}"
        if name == "func":
            return r"\operatorname{" + child("fName") + "} {" + child("e") + "}"
        if name in {"limLow", "limUpp"}:
            command = r"\underset" if name == "limLow" else r"\overset"
            return command + "{" + child("lim") + "}{" + child("e") + "}"
        if name == "bar":
            command = r"\underline" if prop(node, "barPr", "pos", "bot") == "bot" else r"\overline"
            return command + "{" + child("e") + "}"
        if name == "acc":
            accents = {"̂": r"\hat", "̅": r"\bar", "⃗": r"\vec", "̃": r"\tilde", "̇": r"\dot", "̈": r"\ddot"}
            accent = prop(node, "accPr", "chr", "̂")
            if accent not in accents:
                unresolved.add("unhandled_math_accent")
                return r"\overset{" + literal(accent) + "}{" + child("e") + "}"
            return accents[accent] + "{" + child("e") + "}"
        if name not in {"oMath", "oMathPara", "r", "e", "num", "den", "deg", "sub", "sup", "lim", "fName"}:
            unresolved.add("unhandled_omml_" + re.sub(r"[^A-Za-z0-9]", "", name)[:40])
        return joined()

    rendered = render(element)
    if len(rendered) > MAX_EQUATION_CHARS:
        fail("DOCUMENT_LIMIT", "The formula structure is too large. Split the document.")
    return {"format": "latex", "text": rendered, "source_format": "omml", "status": "partial" if unresolved else "complete", "unresolved": sorted(unresolved)}


def word_content(element: Any) -> tuple[str, list[dict[str, Any]]]:
    equations: list[dict[str, Any]] = []

    def visit(node: Any) -> str:
        tag = str(node.tag)
        if tag == f"{{{MATH_NS}}}oMath":
            equation = omml_equation(node)
            equations.append(equation)
            return "\\(" + equation["text"] + "\\)"
        if tag == f"{{{WORD_NS}}}t":
            return node.text or ""
        if tag in {f"{{{WORD_NS}}}br", f"{{{WORD_NS}}}cr"}:
            return "\n"
        if tag == f"{{{WORD_NS}}}tab":
            return "\t"
        if tag in {f"{{{WORD_NS}}}del", f"{{{WORD_NS}}}moveFrom", f"{{{WORD_NS}}}instrText"}:
            return ""  # current visible content, never execute fields
        text = "".join(visit(child) for child in node)
        return text + "\n" if tag == f"{{{WORD_NS}}}p" else text

    return visit(element).rstrip("\n"), equations


class ReaderError(Exception):
    def __init__(self, code: str, message: str):
        self.code, self.message = code, message
        super().__init__(message)


def fail(code: str, message: str) -> None:
    raise ReaderError(code, message)


def bounded_int(value: Any, default: int, low: int, high: int) -> int:
    if value is None:
        return default
    if isinstance(value, bool) or not isinstance(value, int) or not low <= value <= high:
        fail("INVALID_REQUEST", "The reading range is invalid.")
    return value


def _plain_ancestors(path: Path) -> None:
    for item in [path, *path.parents]:
        info = item.lstat()
        if stat.S_ISLNK(info.st_mode) or getattr(info, "st_file_attributes", 0) & 0x400:
            fail("PATH_REPARSE_POINT", "Symbolic links and junctions are not read. Select the actual workspace directory.")


def _inside(child: Path, root: Path) -> bool:
    try:
        child.relative_to(root)
        return True
    except ValueError:
        return False


def _opened_windows_path(handle: Any) -> Path:
    import ctypes
    from ctypes import wintypes
    import msvcrt

    fn = ctypes.windll.kernel32.GetFinalPathNameByHandleW
    fn.argtypes = [wintypes.HANDLE, wintypes.LPWSTR, wintypes.DWORD, wintypes.DWORD]
    fn.restype = wintypes.DWORD
    buffer = ctypes.create_unicode_buffer(32768)
    size = fn(msvcrt.get_osfhandle(handle.fileno()), buffer, len(buffer), 0)
    if not size or size >= len(buffer):
        fail("PATH_UNVERIFIED", "The actual file location could not be verified.")
    name = buffer.value
    if name.startswith("\\\\?\\UNC\\"):
        name = "\\\\" + name[8:]
    elif name.startswith("\\\\?\\"):
        name = name[4:]
    return Path(name)


def source_bytes(request: dict[str, Any]) -> tuple[str, str, bytes, str]:
    root_raw, relative = request.get("root"), request.get("path")
    if not isinstance(root_raw, str) or not isinstance(relative, str) or not relative or len(relative) > 4096:
        fail("INVALID_REQUEST", "An authorized workspace and relative file path are required.")
    parts = PurePosixPath(relative.replace("\\", "/"))
    win = PureWindowsPath(relative)
    if parts.is_absolute() or win.is_absolute() or win.drive or ":" in relative or "\x00" in relative or ".." in parts.parts:
        fail("PATH_OUTSIDE_WORKSPACE", "The file must be inside the authorized workspace.")
    if any(piece.endswith((" ", ".")) for piece in parts.parts):
        fail("INVALID_PATH", "The file path contains an unsupported name.")
    root = Path(root_raw)
    if not root.is_absolute():
        fail("INVALID_REQUEST", "The workspace root must be an absolute path.")
    _plain_ancestors(root)
    root = root.resolve(strict=True)
    if not root.is_dir():
        fail("INVALID_WORKSPACE", "The workspace directory is unavailable.")
    candidate = root.joinpath(*parts.parts)
    _plain_ancestors(candidate)
    candidate = candidate.resolve(strict=True)
    if candidate == root or not _inside(candidate, root):
        fail("PATH_OUTSIDE_WORKSPACE", "The file must be inside the authorized workspace.")
    suffix = candidate.suffix.lower()
    if suffix in {".doc", ".ppt", ".xls"}:
        fail("CONVERSION_REQUIRED", "Save legacy Office binary documents as DOCX, PPTX, or XLSX before reading.")
    if suffix not in SUPPORTED_FORMATS:
        fail("UNSUPPORTED_FORMAT", "Supported formats are PDF, DOCX, PPTX, XLSX, Markdown, plain text, and common images.")
    with candidate.open("rb") as handle:
        before = os.fstat(handle.fileno())
        if not stat.S_ISREG(before.st_mode):
            fail("INVALID_FILE", "Select a regular document file.")
        limit = MAX_TEXT_BYTES if suffix in TEXT_FORMATS else MAX_ASSET_BYTES if suffix in IMAGE_FORMATS else MAX_FILE_BYTES
        if before.st_size > limit:
            fail("FILE_TOO_LARGE", "The document exceeds the local reading size limit. Split it.")
        if os.name == "nt" and _opened_windows_path(handle) != candidate:
            fail("PATH_CHANGED", "The file location changed. Select it again.")
        content = handle.read(limit + 1)
        after = os.fstat(handle.fileno())
        _plain_ancestors(candidate)
        now = candidate.stat()
        if (before.st_ino, before.st_size, before.st_mtime_ns) != (after.st_ino, after.st_size, after.st_mtime_ns) or (after.st_ino, after.st_size, after.st_mtime_ns) != (now.st_ino, now.st_size, now.st_mtime_ns):
            fail("SOURCE_CHANGED", "The document changed during reading. Read it again.")
    if len(content) > limit:
        fail("FILE_TOO_LARGE", "The document exceeds the local reading size limit. Split it.")
    digest = hashlib.sha256(content).hexdigest()
    expected = request.get("document_sha256")
    if expected is not None and expected != digest:
        fail("SOURCE_CHANGED", "The document changed. Read the text again before opening its images.")
    return parts.as_posix(), suffix, content, digest


def validate_zip(content: bytes, suffix: str) -> None:
    if not content.startswith(b"PK\x03\x04"):
        fail("FORMAT_MISMATCH", "The file content does not match its extension.")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        if len(members) > MAX_ZIP_MEMBERS:
            fail("ARCHIVE_LIMIT", "The document has too many internal entries. Split it.")
        names: set[str] = set()
        total = 0
        for info in members:
            path = PurePosixPath(info.filename)
            if info.filename in names or path.is_absolute() or ".." in path.parts or "\\" in info.filename or ":" in info.filename:
                fail("UNSAFE_ARCHIVE", "The document has an unsafe internal structure.")
            names.add(info.filename)
            total += info.file_size
            if info.flag_bits & 1 or stat.S_ISLNK(info.external_attr >> 16):
                fail("UNSAFE_ARCHIVE", "Encrypted Office documents or documents containing links are unsupported.")
            if info.file_size > MAX_ZIP_MEMBER or total > MAX_ZIP_EXPANDED or (info.file_size > 1024 * 1024 and info.file_size / max(1, info.compress_size) > MAX_ZIP_RATIO):
                fail("ARCHIVE_LIMIT", "The expanded document size exceeds the safe reading limit.")
            if info.filename.endswith((".xml", ".rels")):
                xml = archive.read(info)
                # XML permits UTF-16/32 too; removing NULs catches their DTD
                # markers before any third-party parser sees the document.
                declaration_probe = xml.replace(b"\x00", b"").upper()
                if b"<!DOCTYPE" in declaration_probe or b"<!ENTITY" in declaration_probe:
                    fail("UNSAFE_ARCHIVE", "Documents containing external entities are unsupported.")
        if OFFICE_FORMATS[suffix] not in names or "[Content_Types].xml" not in names:
            fail("FORMAT_MISMATCH", "The file content does not match its Office extension.")


def unit(kind: str, text: str, locator: dict[str, Any], assets: list[dict[str, Any]] | None = None, **extra: Any) -> dict[str, Any]:
    if len(text) > MAX_UNIT_CHARS:
        fail("DOCUMENT_LIMIT", "A page or paragraph is too long. Split the document.")
    return {"kind": kind, "text": text, "locator": locator, "assets": assets or [], "extraction": "text", **extra}


def asset_metadata(asset_id: str, locator: dict[str, Any]) -> dict[str, Any]:
    return {"id": asset_id, "mime_type": "image/png", "locator": locator}


class Document:
    def __init__(self, content: bytes, suffix: str, ocr: str, root: str = "", relative: str = ""):
        self.content, self.suffix, self.ocr = content, suffix, ocr
        self.root, self.relative = root, relative
        self.warnings: list[str] = []
        self.assets: dict[str, bytes] = {}
        self.markdown_images: dict[str, dict[str, Any]] = {}
        self.rows: list[dict[str, Any]] = []
        self.pdf: Any = None
        self.ocr_engine: Any = None
        self.ocr_attempts = 0
        self.ocr_deferred = False
        self.unresolved: set[str] = set()
        if suffix == ".pdf":
            if not content.startswith(b"%PDF-"):
                fail("FORMAT_MISMATCH", "The file content does not match its PDF extension.")
            import pymupdf
            pymupdf.TOOLS.mupdf_display_errors(False)
            pymupdf.TOOLS.mupdf_display_warnings(False)
            self.pdf = pymupdf.open(stream=content, filetype="pdf")
            if self.pdf.needs_pass:
                fail("ENCRYPTED_DOCUMENT", "This PDF requires a password. Decrypt a copy before reading.")
            self.count = self.pdf.page_count
        elif suffix in OFFICE_FORMATS:
            validate_zip(content, suffix)
            if suffix == ".docx":
                self.docx()
            elif suffix == ".pptx":
                self.pptx()
            else:
                self.xlsx()
            self.count = len(self.rows)
        elif suffix in IMAGE_FORMATS:
            self.verify_image(content)
            self.assets["image:1"] = content
            self.rows = [unit("image", "", {"image": 1}, [asset_metadata("image:1", {"image": 1})], extraction="none", status="needs_ocr")]
            self.count = 1
        else:
            self.plain_text()
            self.count = len(self.rows)
        if self.count > MAX_UNITS:
            fail("DOCUMENT_LIMIT", "The document has too many content units. Split it.")

    def warn(self, warning: str) -> None:
        if warning not in self.warnings:
            self.warnings.append(warning)

    def append(self, row: dict[str, Any]) -> None:
        self.rows.append(row)
        if len(self.rows) > MAX_UNITS:
            fail("DOCUMENT_LIMIT", "The document has too many content units. Split it.")

    def plain_text(self) -> None:
        raw = self.content
        if raw.startswith((b"%PDF-", b"PK\x03\x04", b"\xd0\xcf\x11\xe0")):
            fail("FORMAT_MISMATCH", "Binary files cannot be read as plain text.")
        encodings = ["utf-8-sig"]
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            encodings.insert(0, "utf-16")
        encodings += ["gb18030"]
        for encoding in encodings:
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            fail("TEXT_ENCODING", "The text encoding is unknown. Save the file as UTF-8 text.")
        if "\x00" in text or sum(ord(c) < 32 and c not in "\r\n\t\f" for c in text) > max(4, len(text) // 100):
            fail("FORMAT_MISMATCH", "The file contains binary content and cannot be read as plain text.")
        # Keep newline characters so cursor reconstruction never joins words.
        for index, line in enumerate(text.splitlines(keepends=True), 1):
            locator = {"line_start": index, "line_end": index}
            row = unit("line", line, locator)
            if self.suffix in {".md", ".markdown"}:
                references = re.findall(r"!\[[^\]\r\n]*\]\(\s*(?:<([^>\r\n]+)>|([^\s)]+))(?:\s+['\"][^\r\n]*?['\"])?\s*\)", line)
                targets = [first or second for first, second in references]
                targets.extend(ref.split("|", 1)[0] for ref in re.findall(r"!\[\[([^\]\r\n]+)\]\]", line))
                row["_image_refs"] = []
                if re.search(r"!\[[^\]\r\n]*\]\s*\[|<img\b", line, re.IGNORECASE):
                    self.warn("Markdown reference images and HTML img elements are not expanded. Use inline relative images or Obsidian image embeds.")
                for number, reference in enumerate(targets[:MAX_IMAGES_PER_UNIT], 1):
                    base_id = f"md-image:{index}:{number}"
                    self.markdown_images[base_id] = {"reference": reference, "locator": locator}
                    row["_image_refs"].append(base_id)
                if len(targets) > MAX_IMAGES_PER_UNIT:
                    self.warn("This line has too many image references. Only the first 32 are shown.")
            self.append(row)

    @staticmethod
    def verify_image(content: bytes) -> None:
        from PIL import Image
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(content)) as image:
                    if image.width * image.height > MAX_IMAGE_PIXELS:
                        fail("IMAGE_TOO_LARGE", "The image exceeds the local pixel limit.")
                    image.verify()
        except ReaderError:
            raise
        except (Image.DecompressionBombWarning, Image.DecompressionBombError):
            fail("IMAGE_TOO_LARGE", "The image exceeds the local pixel limit.")
        except Exception:
            fail("FORMAT_MISMATCH", "The file does not contain a readable image.")

    def markdown_asset(self, base_id: str) -> tuple[bytes, str, dict[str, Any]]:
        info = self.markdown_images.get(base_id)
        if not info:
            fail("ASSET_NOT_FOUND", "The image reference does not exist in the document.")
        reference = unquote(info["reference"]).replace("\\", "/")
        parsed = urlsplit(reference)
        if parsed.scheme or parsed.netloc or reference.startswith("/"):
            fail("EXTERNAL_IMAGE", "Remote image links are not read automatically. Use local images inside the workspace.")
        if parsed.query:
            fail("INVALID_IMAGE_REFERENCE", "The local image reference has unsupported query parameters.")
        relative = posixpath.normpath(posixpath.join(posixpath.dirname(self.relative), parsed.path))
        if PurePosixPath(relative).suffix.lower() not in IMAGE_FORMATS:
            fail("UNSUPPORTED_IMAGE", "This Markdown reference is not a supported local image.")
        _, _, content, digest = source_bytes({"root": self.root, "path": relative})
        self.verify_image(content)
        return content, digest, info["locator"]

    def office_images(self, element: Any, part: Any, prefix: str, locator: dict[str, Any]) -> list[dict[str, Any]]:
        result: list[dict[str, Any]] = []
        relation_attribute = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        relationship_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        seen: set[str] = set()
        for node in element.iter():
            rel_id = node.get(relation_attribute) or node.get(relationship_ns + "link")
            if str(node.tag).endswith("}imagedata"):
                rel_id = rel_id or node.get(relationship_ns + "id")
            if not rel_id:
                continue
            relation = part.rels.get(rel_id)
            if relation is None:
                continue
            if relation.is_external:
                self.warn("The Office document contains externally linked images. No network fetch was made; only embedded images are returned.")
                self.unresolved.add("external_office_image")
                continue
            if rel_id in seen:
                continue
            seen.add(rel_id)
            if len(result) >= MAX_IMAGES_PER_UNIT:
                self.warn("Some image references exceed this unit\'s display limit. All images remain in the original document.")
                self.unresolved.add("image_reference_limit")
                break
            target = relation.target_part
            blob = target.blob
            if len(blob) > MAX_ZIP_MEMBER:
                self.warn("The document contains oversized images that were not loaded into the image API.")
                self.unresolved.add("embedded_image_size_limit")
                continue
            asset_id = f"{prefix}:{rel_id}"
            self.assets[asset_id] = blob
            result.append(asset_metadata(asset_id, locator))
        return result

    def docx(self) -> None:
        from docx import Document as DocxDocument
        from docx.table import Table

        doc = DocxDocument(io.BytesIO(self.content))
        paragraph_number, table_number = 0, 0
        def body_blocks(container: Any):
            for child in container:
                name = str(child.tag).rsplit("}", 1)[-1]
                if name in {"p", "tbl"}:
                    yield child
                elif name not in {"sectPr", "del", "moveFrom", "sdtPr", "sdtEndPr"}:
                    if name not in {"sdt", "sdtContent", "customXml", "ins", "moveTo"}:
                        self.unresolved.add("word_body_container")
                        self.warn("Word contains special content containers that need source verification. Reading is not marked complete.")
                    yield from body_blocks(child)

        for element in body_blocks(doc.element.body):
            table_data = None
            if element.tag.endswith("}p"):
                paragraph_number += 1
                locator = {"paragraph": paragraph_number}
                text, equations = word_content(element)
                kind = "paragraph"
            elif element.tag.endswith("}tbl"):
                table_number += 1
                locator = {"table": table_number}
                table = Table(element, doc)
                table_rows = [[word_content(cell._tc)[0] for cell in row.cells] for row in table.rows]
                text = "\n".join("\t".join(cells) for cells in table_rows)
                _, equations = word_content(element)
                table_data = {"rows": table_rows, "merged_cells": "repeated_in_grid", "source_format": "word_table"}
                kind = "table"
            images = self.office_images(element, doc.part, "docx-image", locator)
            row = unit(kind, text, locator, images, equations=equations)
            if table_data is not None:
                row["table"] = table_data
            unsupported = []
            for node in element.iter():
                name = str(node.tag).rsplit("}", 1)[-1]
                if name in {"OLEObject", "chart", "diagram", "smartTag", "altChunk", "sym"}:
                    unsupported.append("word_object_" + name)
            if unsupported:
                row["unresolved"] = sorted(set(unsupported))
                self.warn("Word contains embedded objects, charts, or special symbols whose structure may not be parsed. Check the source and retained previews.")
            if any(equation["status"] == "partial" for equation in equations):
                self.warn("Some Word formulas retain only readable linear content. Check complex structures in the source.")
            self.append(row)
        # These parts used to be omitted. Keep distinct source locations, not
        # fictitious printed pages, and never execute embedded Office objects.
        seen_parts: set[str] = set()
        for relation in doc.part.rels.values():
            if relation.is_external:
                continue
            part = relation.target_part
            name = str(part.partname)
            if name in seen_parts or not re.fullmatch(r"/word/(?:header\d+|footer\d+|footnotes|endnotes|comments)\.xml", name):
                continue
            seen_parts.add(name)
            from lxml import etree
            tree = etree.fromstring(part.blob, parser=etree.XMLParser(resolve_entities=False, no_network=True))
            locator = {"part": name, "section": "word_supplement"}
            text, equations = word_content(tree)
            images = self.office_images(tree, part, "docx-part-image:" + name.rsplit("/", 1)[-1], locator)
            if text.strip() or images:
                self.append(unit("document_supplement", text, locator, images, equations=equations))
        if any(row.get("equations") for row in self.rows):
            self.warn("Native Word formulas were converted to structured mathematical text. This does not imply recognition of all legacy OLE or image formulas.")

    def pptx(self) -> None:
        from pptx import Presentation

        doc = Presentation(io.BytesIO(self.content))
        for number, slide in enumerate(doc.slides, 1):
            locator = {"slide": number}
            texts: list[str] = []

            def shapes(items: Any) -> None:
                for shape in items:
                    if hasattr(shape, "shapes"):
                        shapes(shape.shapes)
                    elif shape.has_table:
                        texts.append("\n".join("\t".join(cell.text for cell in row.cells) for row in shape.table.rows))
                    elif shape.has_text_frame:
                        texts.append(shape.text_frame.text)

            shapes(slide.shapes)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    texts.append("[Speaker notes]\n" + notes.text)
            images = self.office_images(slide.element, slide.part, f"pptx-image:{number}", locator)
            self.append(unit("slide", "\n".join(texts), locator, images))
        self.warn("PPTX returns text, tables, notes, and embedded images, not a screenshot of the full slide layout.")

    def xlsx(self) -> None:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(self.content), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook:
                if sheet.max_row > MAX_UNITS or sheet.max_column > 256:
                    fail("DOCUMENT_LIMIT", "The worksheet range is too large. Split it or export the selected range.")
                for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    text = "\t".join("" if value is None else str(value) for value in values)
                    self.append(unit("worksheet_row", text, {"sheet": sheet.title, "row": row_number}))
        finally:
            workbook.close()
        self.warn("Worksheet reading returns cell values and formula text. It does not evaluate formulas or read embedded objects.")

    def local_ocr(self, page: Any = None, image_bytes: bytes | None = None) -> str:
        if self.ocr == "off":
            return ""
        if self.ocr_attempts >= 2:
            self.ocr_deferred = True
            self.warn("Scan OCR is limited to two pages per batch. Use the cursor to read other pages or view page images.")
            return ""
        self.ocr_attempts += 1
        try:
            try:
                import rapidocr as ocr_module
                modern = True
                filenames = {"Det.model_path": "PP-OCRv6_det_small.onnx", "Rec.model_path": "PP-OCRv6_rec_small.onnx", "Cls.model_path": "ch_ppocr_mobile_v2.0_cls_mobile.onnx"}
            except ImportError:
                import rapidocr_onnxruntime as ocr_module
                modern = False
                filenames = {"det_model_path": "ch_PP-OCRv3_det_infer.onnx", "rec_model_path": "ch_PP-OCRv3_rec_infer.onnx", "cls_model_path": "ch_ppocr_mobile_v2.0_cls_infer.onnx"}
            package = Path(ocr_module.__file__).resolve().parent
            files = {key: package / "models" / value for key, value in filenames.items()}
            if not all(value.is_file() for value in files.values()):
                self.warn("The local OCR model is unavailable. No model was downloaded; original scan images remain viewable.")
                return ""
            # Optional upstream fallback downloads are disabled even if future
            # dependencies add them. Only bundled, explicit model files run.
            with mock.patch("socket.socket.connect", side_effect=RuntimeError("offline OCR")), mock.patch("socket.socket.connect_ex", side_effect=RuntimeError("offline OCR")), mock.patch("socket.create_connection", side_effect=RuntimeError("offline OCR")):
                if self.ocr_engine is None:
                    paths = {key: str(value) for key, value in files.items()}
                    if modern:
                        self.ocr_engine = ocr_module.RapidOCR(params={**paths, "Global.log_level": "critical", "EngineConfig.onnxruntime.intra_op_num_threads": 2, "EngineConfig.onnxruntime.inter_op_num_threads": 1})
                    else:
                        self.ocr_engine = ocr_module.RapidOCR(print_verbose=False, **paths)
                pixels = self.pdf_pixels(page, 1600) if image_bytes is None else image_bytes
                results = self.ocr_engine(pixels)
            if modern:
                return "\n".join(str(text) for text in getattr(results, "txts", None) or [])
            return "\n".join(str(item[1]) for item in results[0] or [] if len(item) >= 2)
        except Exception:
            self.warn("Local OCR is unavailable or recognition failed. No model was downloaded; original scan images remain viewable.")
            return ""

    @staticmethod
    def pdf_pixels(page: Any, max_dimension: int) -> bytes:
        import pymupdf
        edge = max(page.rect.width, page.rect.height)
        if edge <= 0:
            fail("CORRUPT_DOCUMENT", "The PDF page dimensions are invalid.")
        scale = min(2.0, max_dimension / edge)
        image = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
        return image.tobytes("png")

    def get(self, index: int) -> dict[str, Any]:
        if self.pdf is None:
            row = dict(self.rows[index])
            row["assets"] = list(row["assets"])
            for base_id in row.pop("_image_refs", []):
                try:
                    _, digest, locator = self.markdown_asset(base_id)
                    metadata = asset_metadata(f"{base_id}:{digest}", locator)
                    metadata["sha256"] = digest
                    row["assets"].append(metadata)
                except ReaderError as exc:
                    row.setdefault("unresolved", []).append("image_reference_unavailable")
                    self.warn(f"Markdown image not read: {exc.message}")
                except (OSError, ValueError):
                    row.setdefault("unresolved", []).append("image_reference_unavailable")
                    self.warn("A local Markdown image is unavailable. Check its file and relative path.")
            if self.suffix in IMAGE_FORMATS or self.suffix == ".docx":
                additions = []
                for asset in row["assets"]:
                    try:
                        normalized = self.asset(asset["id"], 1600)
                        recognized = self.local_ocr(image_bytes=base64.b64decode(normalized["base64"]))
                    except ReaderError:
                        row.setdefault("unresolved", []).append("embedded_image_unreadable")
                        self.warn("Some embedded images cannot be displayed locally. Check the source document.")
                        continue
                    if recognized.strip():
                        additions.append(recognized)
                    else:
                        row.setdefault("unresolved", []).append("embedded_image_text_unverified")
                        row["status"] = "partial" if row["text"].strip() else "needs_ocr"
                        self.warn("Some embedded images have no reliable OCR text. View the original; body text does not prove the image was fully read.")
                if additions:
                    had_text = bool(row["text"].strip())
                    row["text"] += ("\n[Embedded image OCR; verify against the original]\n" if had_text else "") + "\n".join(additions)
                    row["extraction"] = "text+ocr" if had_text else "ocr"
                    row["ocr_applied"] = True
                    if "embedded_image_text_unverified" in row.get("unresolved", []):
                        row["status"] = "partial"
                    else:
                        row.pop("status", None)
                    self.warn("OCR text may contain recognition errors. Verify formulas, charts, and figures against original images.")
            return self.with_coverage(row)
        page = self.pdf[index]
        locator = {"page": index + 1, "page_kind": "pdf_physical"}
        text = page.get_text("text", sort=True) or ""
        assets = [asset_metadata(f"pdf-page:{index + 1}", locator)]
        result = unit("page", text, locator, assets)
        has_images = bool(page.get_images(full=False))
        has_visual = has_images or bool(page.get_drawings())
        # Native PDF text may coexist with an image-only scan. A page number
        # or header must not prevent OCR of the substantive scanned content.
        if has_images or (not text.strip() and has_visual):
            recognized = self.local_ocr(page)
            if len(recognized) > MAX_UNIT_CHARS:
                fail("DOCUMENT_LIMIT", "The recognized scan text is too long. Split the document.")
            if recognized.strip():
                result["text"] = text + "\n[Page image OCR; may duplicate body text]\n" + recognized if text.strip() else recognized
                result["extraction"] = "text+ocr" if text.strip() else "ocr"
                result["ocr_applied"] = True
                self.warn("OCR text may contain recognition errors. Check the corresponding physical page image.")
            else:
                result["unresolved"] = ["raster_or_vector_text_unverified"]
                result["status"] = "partial" if text.strip() else "needs_ocr"
                if not text.strip():
                    result["extraction"] = "none"
        elif not text.strip():
            result["status"] = "empty"
            result["extraction"] = "none"
        if has_visual or text.strip():
            self.warn("Original images are retained for every read physical PDF page. Text extraction or basic OCR does not guarantee formula, table, chart, or reading-order reconstruction.")
            result["visual_review_required"] = True
        return self.with_coverage(result)

    def with_coverage(self, row: dict[str, Any]) -> dict[str, Any]:
        if len(row["text"]) > MAX_UNIT_CHARS:
            fail("DOCUMENT_LIMIT", "A page or paragraph is too long. Split the document.")
        unresolved = set(row.get("unresolved", []))
        for equation in row.get("equations", []):
            unresolved.update(equation.get("unresolved", []))
        if row.get("status") == "needs_ocr":
            unresolved.add("text_not_recognized")
        row["coverage"] = {
            "text_method": row.get("extraction", "text"),
            "visual_assets": len(row["assets"]),
            "formulas_extracted": len(row.get("equations", [])),
            "table_structure": "extracted" if row.get("table") else "not_evaluated",
            "visual_review_required": bool(row.get("visual_review_required") or row["assets"] or unresolved),
            "unresolved": sorted(unresolved),
        }
        if not row["text"].strip() and row["assets"] and row.get("status") != "empty":
            row["status"] = "needs_ocr"
            row["coverage"]["unresolved"] = sorted(unresolved | {"text_not_recognized"})
        return row

    def asset(self, asset_id: str, max_dimension: int) -> dict[str, Any]:
        locator: dict[str, Any]
        if self.pdf is not None:
            match = re.fullmatch(r"pdf-page:([1-9][0-9]{0,5})", asset_id)
            if not match or not 1 <= int(match[1]) <= self.count:
                fail("ASSET_NOT_FOUND", "The page image does not exist in this document.")
            number = int(match[1])
            data = self.pdf_pixels(self.pdf[number - 1], max_dimension)
            locator = {"page": number, "page_kind": "pdf_physical"}
        elif asset_id.startswith("md-image:"):
            if not re.fullmatch(r"md-image:[1-9][0-9]*:[1-9][0-9]*:[0-9a-f]{64}", asset_id):
                fail("ASSET_NOT_FOUND", "The document image version identifier is invalid.")
            base_id, expected = asset_id.rsplit(":", 1)
            data, digest, locator = self.markdown_asset(base_id)
            if digest != expected:
                fail("SOURCE_CHANGED", "The referenced Markdown image changed. Read the corresponding source chunk again.")
        else:
            if asset_id not in self.assets:
                fail("ASSET_NOT_FOUND", "The image does not exist in this document.")
            data = self.assets[asset_id]
            locator = next(asset["locator"] for row in self.rows for asset in row["assets"] if asset["id"] == asset_id)
        from PIL import Image
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            try:
                with Image.open(io.BytesIO(data)) as image:
                    if image.width * image.height > MAX_IMAGE_PIXELS:
                        fail("IMAGE_TOO_LARGE", "The image exceeds the local pixel limit.")
                    image.thumbnail((max_dimension, max_dimension))
                    if image.mode not in {"RGB", "RGBA"}:
                        image = image.convert("RGBA" if "transparency" in image.info else "RGB")
                    output = io.BytesIO()
                    image.save(output, format="PNG")
                    png = output.getvalue()
                    width, height = image.size
            except (Image.DecompressionBombWarning, Image.DecompressionBombError):
                fail("IMAGE_TOO_LARGE", "The image exceeds the local pixel limit.")
            except ReaderError:
                raise
            except Exception:
                fail("UNSUPPORTED_IMAGE", "This embedded image format cannot be converted locally. View it in the source.")
        if len(png) > MAX_ASSET_BYTES:
            fail("IMAGE_TOO_LARGE", "The converted image is too large. Reduce its dimensions and retry.")
        return {"id": asset_id, "mime_type": "image/png", "base64": base64.b64encode(png).decode("ascii"), "width": width, "height": height, "locator": locator}

    def close(self) -> None:
        if self.pdf is not None:
            self.pdf.close()


# Placeholder values that models and form fields send for "no value". They mean
# "start from the beginning", never a position, so they must not fail a read.
NULL_TOKENS = {"", "null", "none", "undefined", "nil", "start", "begin", "beginning", "first", "0"}
_POSITION_TEXT = re.compile(r"(?:start_unit|unit|page|slide|line|p)\s*[:=#]?\s*(\d{1,9})", re.IGNORECASE)
_INT_LIMITS = {"start_unit": (1, MAX_UNITS), "max_chars": (256, 100000), "max_units": (1, 40), "max_dimension": (128, 2048)}
_ENGINE_ALIASES = {"auto": "auto", "default": "auto", "builtin": "builtin", "native": "builtin", "fast": "builtin",
                   "quick": "builtin", "text": "builtin", "pymupdf": "builtin", "local": "builtin",
                   "mineru": "mineru", "structure": "mineru", "structured": "mineru", "advanced": "mineru", "layout": "mineru"}
# MinerU tiers (verified in MinerU 4.0.12 parser/tier.py and backend/analysis):
# basic = ONNX small models only, no VLM; standard = hybrid layout + local VLM;
# advanced = full local VLM two-step extraction with image analysis.
# adaptive (DeepSteam default) = basic first, standard only for pages whose
# basic layout shows tables or display formulas.
_TIER_ALIASES = {"adaptive": "adaptive", "auto": "adaptive", "smart": "adaptive", "mixed": "adaptive",
                 "basic": "basic", "pipeline": "basic", "flash": "basic", "fast": "basic", "low": "basic",
                 "standard": "standard", "hybrid": "standard", "medium": "standard",
                 "advanced": "advanced", "vlm": "advanced", "high": "advanced", "pro": "advanced",
                 "full": "advanced", "enhanced": "advanced", "best": "advanced", "xhigh": "advanced"}


def normalize_cursor(value: Any) -> str | None:
    """Canonical cursor token, or None for "read from the beginning".

    Tolerates the transport damage models commonly introduce when copying an
    opaque token: surrounding quotes/whitespace, dropped ``=`` padding and the
    standard (non-URL) base64 alphabet. The decoded content is still bound to
    the document SHA-256, so a repaired token cannot address another source.
    """
    if value is None:
        return None
    if not isinstance(value, str):
        fail("INVALID_CURSOR", "The pagination cursor must be the exact next_cursor text. Omit it to start at the first unit, or pass start_unit to continue at a page.")
    token = value.strip().strip("\"'` ")
    if token.lower() in NULL_TOKENS:
        return None
    token = token.replace("+", "-").replace("/", "_").rstrip("=")
    if len(token) > 1024 or not re.fullmatch(r"[A-Za-z0-9_-]+", token):
        fail("INVALID_CURSOR", "The pagination cursor is damaged. Omit it to start at the first unit, or pass start_unit (one-based page/slide/line) to continue.")
    return token + "=" * (-len(token) % 4)


def cursor_payload(value: Any) -> dict[str, Any] | None:
    """Decoded cursor object without the source check, or None if undecodable."""
    try:
        token = normalize_cursor(value)
        if token is None:
            return None
        cursor = json.loads(base64.urlsafe_b64decode(token))
        if not isinstance(cursor, dict) or cursor.get("v") != SCHEMA_VERSION or type(cursor.get("index")) is not int:
            return None
        return cursor
    except (ReaderError, ValueError, TypeError):
        return None


def normalize_read_request(request: dict[str, Any]) -> dict[str, Any]:
    """Repair harmless caller mistakes before they become reading failures.

    Empty placeholders ("", null, 0) become defaults and numeric strings become
    integers; explicit out-of-range values are still rejected with the allowed
    range. A plain page number sent as a cursor becomes ``start_unit``, and a
    valid cursor wins over a conflicting ``start_unit``. Unknown engines, tiers
    and OCR modes (for example "cloud") are rejected: everything runs locally.
    """
    value = dict(request)
    ocr = value.get("ocr")
    if isinstance(ocr, bool):
        value["ocr"] = "auto" if ocr else "off"
    elif isinstance(ocr, str):
        flag = ocr.strip().lower()
        if flag in {"", "null", "undefined", "default", "auto", "on", "true", "yes", "1", "enabled", "local"}:
            value["ocr"] = "auto"
        elif flag in {"off", "false", "no", "0", "disabled", "none"}:
            value["ocr"] = "off"
        # Anything else (for example "cloud") stays as sent and is rejected:
        # OCR only ever runs locally.
    elif "ocr" in value and ocr is None:
        value.pop("ocr")
    for key in ("cursor", "job_id", "engine", "tier", "asset_id", "document_sha256", "coverage_cursor"):
        item = value.get(key)
        if key in value and (item is None or isinstance(item, str) and item.strip().strip("\"'` ").lower() in NULL_TOKENS):
            value.pop(key)
        elif isinstance(item, str) and key in {"job_id", "document_sha256", "asset_id"}:
            value[key] = item.strip()
    for key, (low, high) in _INT_LIMITS.items():
        item = value.get(key)
        if isinstance(item, str) and re.fullmatch(r"\s*\d{1,9}\s*", item):
            item = int(item)
        elif isinstance(item, float) and item.is_integer():
            item = int(item)
        if key not in value:
            continue
        if item is None or item == "" or item == 0 or (isinstance(item, str) and item.strip().lower() in NULL_TOKENS):
            value.pop(key)  # A blank placeholder means the documented default.
        elif isinstance(item, bool) or not isinstance(item, int) or item < low or item > high:
            fail("INVALID_REQUEST", f"{key} must be an integer from {low} to {high}; omit it to use the default.")
        else:
            value[key] = item
    cursor = value.get("cursor")
    if isinstance(cursor, str):
        text = cursor.strip().strip("\"'` ")
        position = re.fullmatch(r"\d{1,9}", text) or _POSITION_TEXT.fullmatch(text)
        if position:
            # A page number is not an opaque cursor. Reading from that unit can
            # repeat text already seen but can never skip unread content.
            value.pop("cursor")
            number = int(position.group(1) if position.lastindex else position.group(0))
            if number >= 1 and "start_unit" not in value:
                value["start_unit"] = min(number, MAX_UNITS)
        elif "start_unit" in value:
            if cursor_payload(cursor) is not None:
                value.pop("start_unit")
            else:
                value.pop("cursor")
    if isinstance(value.get("cursor"), str):
        try:
            value["cursor"] = normalize_cursor(value["cursor"])
        except ReaderError:
            pass  # Reported with recovery guidance when the cursor is used.
        if value["cursor"] is None:
            value.pop("cursor")
    if isinstance(value.get("engine"), str):
        engine = _ENGINE_ALIASES.get(value["engine"].strip().lower())
        if engine is None:
            fail("INVALID_REQUEST", "engine must be auto, builtin or mineru. All engines run locally; no parsing API is used.")
        value["engine"] = engine
    if isinstance(value.get("tier"), str):
        tier = _TIER_ALIASES.get(value["tier"].strip().lower())
        if tier is None and value["tier"].strip().lower() != "default":
            fail("INVALID_REQUEST", "tier must be adaptive (default: small local models, local VLM only for tables and formulas), standard (hybrid local VLM), advanced (local VLM with image analysis) or basic (small local models, no VLM).")
        if tier is None:
            value.pop("tier")
        else:
            value["tier"] = tier
    return value


def cursor_decode(value: Any, digest: str) -> tuple[int, int]:
    token = normalize_cursor(value)
    if token is None:
        return 0, 0
    try:
        cursor = json.loads(base64.urlsafe_b64decode(token))
        if not isinstance(cursor, dict) or cursor.get("v") != SCHEMA_VERSION or not isinstance(cursor.get("sha256"), str):
            raise ValueError
        if cursor["sha256"] != digest:
            fail("SOURCE_CHANGED", "The document changed during pagination. Read it again from the beginning.")
        return bounded_int(cursor["index"], 0, 0, MAX_UNITS), bounded_int(cursor["offset"], 0, 0, MAX_UNIT_CHARS)
    except ReaderError:
        raise
    except Exception:
        fail("INVALID_CURSOR", "The pagination cursor is not a next_cursor from this document. Omit it to start at the first unit, or pass start_unit (one-based page/slide/line) to continue.")


def navigation(document_count: int, rows: list[dict[str, Any]], truncated: bool, index: int, offset: int, cursor: str | None) -> dict[str, Any]:
    """Model-friendly position summary. ``next_start_unit`` is a plain integer
    fallback for ``next_cursor``: it restarts the unit containing the cursor,
    so it may repeat part of that unit but never skips unread text."""
    returned = [row.get("unit_index") for row in rows if isinstance(row.get("unit_index"), int)]
    return {"unit_count": document_count,
            "returned_units": [returned[0], returned[-1]] if returned else [],
            "complete_through_unit": index if truncated else document_count,
            "next_cursor": cursor,
            "next_start_unit": index + 1 if truncated else None,
            "next_resumes_mid_unit": bool(truncated and offset),
            "done": not truncated}


def cursor_encode(digest: str, index: int, offset: int) -> str:
    return base64.urlsafe_b64encode(json.dumps({"v": SCHEMA_VERSION, "sha256": digest, "index": index, "offset": offset}, separators=(",", ":")).encode()).decode()


def handle_request(request: Any) -> dict[str, Any]:
    document: Document | None = None
    try:
        if not isinstance(request, dict):
            fail("INVALID_REQUEST", "The request must be a JSON object.")
        request = normalize_read_request(request)
        action = request.get("action", "read")
        if action not in {"read", "asset"}:
            fail("INVALID_REQUEST", "This document operation is unsupported.")
        ocr = request.get("ocr", "auto")
        if ocr not in {"auto", "off"}:
            fail("INVALID_REQUEST", "ocr must be auto or off. OCR always runs locally; no OCR service is used.")
        if action == "asset" and not re.fullmatch(r"[0-9a-f]{64}", str(request.get("document_sha256", ""))):
            fail("SOURCE_VERSION_REQUIRED", "A document reading version is required before opening images.")
        relative, suffix, content, digest = source_bytes(request)
        document = Document(content, suffix, ocr if action == "read" else "off", request["root"], relative)
        common = {"schema_version": SCHEMA_VERSION, "ok": True, "path": relative, "format": suffix.lstrip("."), "document": {"sha256": digest, "bytes": len(content), "unit_count": document.count}, "document_sha256": digest}
        if action == "asset":
            asset_id = request.get("asset_id")
            if not isinstance(asset_id, str) or len(asset_id) > 256:
                fail("INVALID_REQUEST", "The image identifier is invalid.")
            asset = document.asset(asset_id, bounded_int(request.get("max_dimension"), 1600, 128, 2048))
            return {**common, "status": "ready", "asset": asset, "warnings": document.warnings}
        max_chars = bounded_int(request.get("max_chars"), 12000, 256, 100000)
        max_units = bounded_int(request.get("max_units"), 10, 1, 40)
        index, offset = cursor_decode(request.get("cursor"), digest)
        if request.get("start_unit") is not None:
            if request.get("cursor") is not None:
                fail("INVALID_REQUEST", "A starting position and pagination cursor cannot both be specified.")
            index = bounded_int(request["start_unit"], 1, 1, MAX_UNITS) - 1
        if index > document.count or index == document.count and offset:
            fail("POSITION_OUT_OF_RANGE", f"The requested position is after the last unit. This document has {document.count} units; use start_unit between 1 and {max(1, document.count)}.")
        begins_document = index == 0 and offset == 0
        rows: list[dict[str, Any]] = []
        remaining = max_chars
        while index < document.count and len(rows) < max_units and remaining > 0:
            if rows:
                remaining -= 2  # Joining units must also respect max_chars.
                if remaining <= 0:
                    break
            document.ocr_deferred = False
            row = dict(document.get(index))
            # Do not advance past a page whose OCR was skipped only because
            # this request used its budget. The next cursor starts here and
            # gets a fresh budget. A single rich Office unit can itself have
            # more images than the budget; retain that unit as explicit partial.
            if document.ocr_deferred and rows:
                break
            text = row["text"]
            row["unit_index"] = index + 1
            row["unit_text_sha256"] = hashlib.sha256(text.encode('utf-8')).hexdigest()
            if offset > len(text):
                # The unit text can differ between requests (for example a page
                # read with OCR, then continued with ocr=off). Restart this unit
                # instead of failing: text may repeat, but none is skipped.
                document.warn(f"Unit {index + 1} changed length since the previous page of results, so it is returned again from its beginning.")
                offset = 0
                begins_document = index == 0
            end = min(len(text), offset + remaining)
            row["text"] = text[offset:end]
            row["char_start"], row["char_end"], row["total_chars"] = offset, end, len(text)
            row["truncated_unit"] = end < len(text)
            if offset or end < len(text):
                # Full table/equation metadata must not bypass a bounded text
                # request. Their linear source text remains losslessly paged.
                row.pop("table", None)
                row.pop("equations", None)
                row["coverage"] = {**row.get("coverage", {}), "structured_content": "fragmented", "formulas_extracted": 0, "table_structure": "fragmented"}
            rows.append(row)
            remaining -= end - offset
            if end < len(text):
                offset = end
                break
            index, offset = index + 1, 0
        truncated = index < document.count
        has_text = any(row["text"].strip() for row in rows)
        needs_ocr = any(row.get("status") == "needs_ocr" or (not row["text"].strip() and row["assets"] and row.get("status") != "empty") for row in rows)
        unresolved = document.unresolved | {code for row in rows for code in row.get("coverage", {}).get("unresolved", [])}
        status = "partial" if has_text and (truncated or needs_ocr or unresolved) else "ready" if has_text else "needs_ocr" if needs_ocr else "partial" if unresolved else "empty"
        if needs_ocr:
            document.warn("Some content has no extractable text. View its images or enable available local OCR.")
        coverage = {
            "scope": "returned_units", "document_fully_read": begins_document and not truncated,
            "semantic_completeness": "not_guaranteed", "returned_unit_count": len(rows),
            "text_units": sum(bool(row["text"].strip()) for row in rows),
            "ocr_units": sum(row.get("extraction") in {"ocr", "text+ocr"} for row in rows),
            "visual_units": sum(bool(row["assets"]) for row in rows),
            "formula_count": sum(len(row.get("equations", [])) for row in rows),
            "table_count": sum(bool(row.get("table")) for row in rows),
            "visual_review_required": any(row.get("coverage", {}).get("visual_review_required", False) for row in rows),
            "unresolved": sorted(unresolved),
            "returned_locators": [row.get("locator", {}) for row in rows],
            "pending_locators": [row.get("locator", {}) for row in rows if row.get("status") in {"partial", "needs_ocr"}],
            "text_extraction_complete": not truncated and not needs_ocr and not unresolved,
        }
        next_cursor = cursor_encode(digest, index, offset) if truncated else None
        return {**common, "status": status, "navigation": navigation(document.count, rows, truncated, index, offset, next_cursor), "text": "\n\n".join(row["text"] for row in rows), "units": rows, "truncated": truncated, "next_cursor": next_cursor, "next_start_unit": index + 1 if truncated else None, "warnings": document.warnings, "coverage": coverage}
    except ReaderError as exc:
        return {"schema_version": SCHEMA_VERSION, "ok": False, "status": "error", "error": {"code": exc.code, "message": exc.message}}
    except FileNotFoundError:
        return {"schema_version": SCHEMA_VERSION, "ok": False, "status": "error", "error": {"code": "FILE_NOT_FOUND", "message": "The workspace document does not exist or was moved."}}
    except PermissionError:
        return {"schema_version": SCHEMA_VERSION, "ok": False, "status": "error", "error": {"code": "FILE_NOT_READABLE", "message": "The document cannot be read. Check permissions and file locks."}}
    except ImportError:
        return {"schema_version": SCHEMA_VERSION, "ok": False, "status": "error", "error": {"code": "PARSER_UNAVAILABLE", "message": "The local document reader is missing. Repair the DeepSteam runtime."}}
    except Exception:
        return {"schema_version": SCHEMA_VERSION, "ok": False, "status": "error", "error": {"code": "CORRUPT_DOCUMENT", "message": "The document could not be parsed. Check its format and integrity."}}
    finally:
        if document is not None:
            document.close()


def main() -> int:
    # Library diagnostics must not leak filenames/content or corrupt JSON.
    logging.disable(logging.CRITICAL)
    request_bytes = sys.stdin.buffer.read(65537)
    if len(request_bytes) > 65536:
        response = {"ok": False, "status": "error", "error": {"code": "INVALID_REQUEST", "message": "The reading request is too large."}}
    else:
        try:
            request = json.loads(request_bytes.decode("utf-8-sig"))
        except Exception:
            request = None
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            response = handle_request(request)
    sys.stdout.buffer.write(json.dumps(response, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
    sys.stdout.buffer.flush()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
