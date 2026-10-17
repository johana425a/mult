"""Nonempty, exact-case artifact acceptance. Local bytes are reopened by format."""
import hashlib
import json
from pathlib import Path


def inspect(path, expected_hash, expected_bytes):
    path = Path(path)
    data = path.read_bytes()
    if not data or len(data) != expected_bytes or hashlib.sha256(data).hexdigest() != expected_hash:
        raise ValueError('Artifact byte/hash mismatch: ' + path.name)
    suffix = path.suffix.lower()
    if suffix == '.docx':
        from docx import Document
        doc = Document(path)
        text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    elif suffix == '.pdf':
        import fitz
        with fitz.open(path) as doc:
            if doc.page_count < 1:
                raise ValueError('PDF has no pages')
            text = '\n'.join(page.get_text() for page in doc)
    elif suffix == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True)
        text = '\n'.join(str(value) for sheet in book for row in sheet.values for value in row if value is not None)
        book.close()
    elif suffix == '.pptx':
        from pptx import Presentation
        text = '\n'.join(shape.text for slide in Presentation(path).slides for shape in slide.shapes if shape.has_text_frame)
    elif suffix in {'.md', '.txt', '.json'}:
        text = data.decode('utf-8-sig')
        if suffix == '.json':
            json.loads(text)
    else:
        raise ValueError('No parser for the requested acceptance format')
    if not text.strip():
        raise ValueError('Parsed artifact has no text')
    return text


def verify(manifest):
    expected = manifest.get('expected_cases')
    cases = manifest.get('cases')
    if not isinstance(expected, dict) or not expected or not isinstance(cases, list) or not cases:
        raise ValueError('Expected acceptance cases and actual results must be nonempty')
    ids = [case['id'] for case in cases]
    if len(ids) != len(set(ids)) or set(ids) != set(expected):
        raise ValueError('Missing, duplicate, or unexpected acceptance case')
    checked = []
    for case in cases:
        files = case.get('files') or []
        formats = [Path(row['path']).suffix.lower().lstrip('.') for row in files]
        if not files or not set(expected[case['id']]).issubset(formats):
            raise ValueError('Required formats are missing from ' + case['id'])
        for row in files:
            text = inspect(row['path'], row['sha256'], row['bytes'])
            markers = row.get('expected_text')
            if not markers or not all(marker in text for marker in markers):
                raise ValueError('Expected source content missing from ' + Path(row['path']).name)
            if row.get('download_path'):
                inspect(row['download_path'], row['sha256'], row['bytes'])
            elif manifest.get('require_downloads', True):
                raise ValueError('Actual downloaded bytes not supplied')
            checked.append({'case': case['id'], 'filename': Path(row['path']).name, 'sha256': row['sha256'], 'bytes': row['bytes'], 'parsed': True})
    return {'passed': True, 'case_count': len(cases), 'file_count': len(checked), 'files': checked}


if __name__ == '__main__':
    import sys
    print(json.dumps(verify(json.loads(Path(sys.argv[1]).read_text('utf-8'))), ensure_ascii=False, indent=2))
