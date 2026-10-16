"""Offline portable-document self-check using only generated, non-user files.

Run with this installation's Python using -I -B. This checks real imports,
text-layer extraction and original image retrieval, not OCR or VLM inference.
It does not start
DeepSteam, read accounts, call a language model, or access a user's workspace.
"""
from __future__ import annotations

import argparse
import base64
import contextlib
import importlib
import importlib.metadata
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import time
from unittest import mock


IMPORTS = {
    "pymupdf": "pymupdf", "python-docx": "docx", "python-pptx": "pptx",
    "openpyxl": "openpyxl", "pillow": "PIL", "lxml": "lxml.etree",
    "defusedxml": "defusedxml", "rapidocr": "rapidocr", "onnxruntime": "onnxruntime",
    "opencv-python-headless": "cv2", "pyclipper": "pyclipper", "Shapely": "shapely",
    "omegaconf": "omegaconf", "colorlog": "colorlog", "flatbuffers": "flatbuffers",
    "protobuf": "google.protobuf",
}


def read_pins(root: Path) -> dict[str, str]:
    pins = {}
    for line in (root / "installer/document-runtime-requirements.txt").read_text("utf8").splitlines():
        if line.strip() and not line.startswith("#"):
            name, version = line.strip().split("==")
            pins[name] = version
    if set(pins) != set(IMPORTS):
        raise RuntimeError("Document runtime pin/import catalogs differ")
    return pins


def verify(root: Path, work_directory: Path) -> dict:
    expected_python = (root / "runtime/python/python.exe").resolve()
    if (not sys.flags.isolated or sys.version_info[:2] != (3, 13)
            or Path(sys.executable).resolve() != expected_python):
        raise RuntimeError("Use this package's portable Python 3.13 with -I -B")
    site = (root / "runtime/python/Lib/site-packages").resolve()
    imports, cases, attempts = {}, [], []

    def block_network(*args, **kwargs):
        attempts.append(True)
        raise RuntimeError("The document self-check is strictly offline")

    started = time.monotonic()
    with contextlib.ExitStack() as stack:
        for target in ("socket.socket.connect", "socket.socket.connect_ex", "socket.getaddrinfo", "urllib.request.urlopen"):
            stack.enter_context(mock.patch(target, side_effect=block_network))
        for distribution, version in read_pins(root).items():
            actual = importlib.metadata.version(distribution)
            if actual != version:
                raise RuntimeError(f"Document component version differs: {distribution}")
            imported = importlib.import_module(IMPORTS[distribution])
            origin = Path(imported.__file__).resolve()
            origin.relative_to(site)
            imports[distribution] = {"version": actual, "origin": origin.relative_to(root).as_posix()}
        reader_path = root / "scripts/workspace_document_reader.py"
        spec = importlib.util.spec_from_file_location("deepsteam_reader_self_check", reader_path)
        reader = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(reader)
        from PIL import Image, ImageDraw, ImageFont
        import pymupdf
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        from openpyxl import Workbook
        from rapidocr.utils.download_file import DownloadFile
        stack.enter_context(mock.patch.object(DownloadFile, "run", side_effect=block_network))
        with tempfile.TemporaryDirectory(prefix="document-self-check-", dir=work_directory) as temp:
            folder = Path(temp)
            picture = Image.new("RGB", (1200, 350), "white")
            ImageDraw.Draw(picture).text((45, 105), "DeepSteam document reading", font=ImageFont.load_default(size=46), fill="black")
            picture.save(folder / "sample.png")
            buffer = io.BytesIO()
            picture.save(buffer, format="PNG")
            with pymupdf.open() as pdf:
                pdf.new_page().insert_text((60, 75), "DeepSteam text PDF")
                pdf.save(folder / "text.pdf")
            with pymupdf.open() as pdf:
                page = pdf.new_page(width=1200, height=350)
                page.insert_image(page.rect, stream=buffer.getvalue())
                pdf.save(folder / "scan.pdf")
            # The generated English scan fixture is packaged so acceptance
            # does not depend on fonts installed on another computer.
            with pymupdf.open() as pdf:
                page = pdf.new_page(width=1200, height=350)
                page.insert_image(page.rect, filename=str(root / "installer/document-fixtures/scanned-en.png"))
                pdf.save(folder / "scan-en.pdf")
            word = Document()
            word.add_paragraph("DeepSteam Word document")
            table = word.add_table(rows=1, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = "Header", "42"
            word.add_picture(str(folder / "sample.png"))
            word.save(folder / "sample.docx")
            slides = Presentation()
            slides.slides.add_slide(slides.slide_layouts[6])
            slide = slides.slides.add_slide(slides.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "DeepSteam slide two"
            slide.shapes.add_picture(str(folder / "sample.png"), Inches(1), Inches(2), width=Inches(6))
            slides.save(folder / "sample.pptx")
            workbook = Workbook()
            workbook.active.append(["DeepSteam spreadsheet", 2, "=B1+1"])
            workbook.save(folder / "sample.xlsx")
            (folder / "sample.md").write_text("# DeepSteam Markdown\n![source](sample.png)\n", "utf8")
            samples = [("text.pdf", ["DeepSteam", "text PDF"]), ("scan.pdf", []),
                       ("scan-en.pdf", []),
                       ("sample.docx", ["Word document", "Header", "42"]), ("sample.pptx", ["slide two"]),
                       ("sample.xlsx", ["spreadsheet", "=B1+1"]), ("sample.md", ["Markdown"])]
            for filename, terms in samples:
                value = reader.handle_request({"root": str(folder), "path": filename, "max_units": 20, "ocr": "off"})
                if not value.get("ok") or not all(term in value.get("text", "") for term in terms):
                    raise RuntimeError(f"Synthetic document failed: {filename}")
                if filename in {"scan.pdf", "scan-en.pdf"}:
                    if value.get('text', '').strip() or not value.get('visual_pending'):
                        raise RuntimeError("Scanned PDF must retain pixels as visually pending, without invented OCR")
                if any(unit.get('extraction') == 'ocr' for unit in value.get('units', [])):
                    raise RuntimeError("Default document self-check unexpectedly invoked OCR")
                asset_count = sum(len(unit.get("assets", [])) for unit in value["units"])
                if filename != "sample.xlsx" and not asset_count:
                    raise RuntimeError(f"Synthetic source visual missing: {filename}")
                cases.append({"file": filename, "passed": True, "status": value["status"],
                              "assets": asset_count, "expected_terms_found": terms})
            # Exercise actual source-page image retrieval as well as metadata.
            pdf_read = reader.handle_request({"root": str(folder), "path": "scan.pdf", "max_units": 1, "ocr": "off"})
            asset = pdf_read["units"][0]["assets"][0]
            visual = reader.handle_request({"action": "asset", "root": str(folder), "path": "scan.pdf",
                                           "asset_id": asset["id"], "document_sha256": pdf_read["document"]["sha256"]})
            if not visual.get("ok") or not base64.b64decode(visual["asset"]["base64"]).startswith(b"\x89PNG\r\n\x1a\n"):
                raise RuntimeError("Original PDF page image could not be read")
    if attempts:
        raise RuntimeError("Document self-check attempted a network request")
    return {"schema": "deepsteam-document-runtime-check-v1", "passed": True, "fixtures_only": True,
            "portable_python": sys.version.split()[0], "python_isolated": True, "imports": imports,
            "cases": cases, "source_image_read_passed": True, "network_attempts": len(attempts),
            "model_requests": 0, "elapsed_seconds": round(time.monotonic() - started, 3)}


def main():
    # Installer pipes must use the same encoding on every Windows locale.
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--work-directory", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    root = args.root.resolve()
    work = (args.work_directory or root / "runtime").resolve()
    if not work.is_dir():
        parser.error("Self-check work directory must already exist")
    try:
        output = verify(root, work)
    except Exception as error:
        output = {"schema": "deepsteam-document-runtime-check-v1", "passed": False,
                  "error_type": type(error).__name__, "error": str(error), "model_requests": 0}
    serialized = json.dumps(output, ensure_ascii=False, indent=2)
    if args.report:
        with args.report.open("x", encoding="utf8") as stream:
            stream.write(serialized)
    print(serialized)
    raise SystemExit(0 if output["passed"] else 1)


if __name__ == "__main__":
    main()
