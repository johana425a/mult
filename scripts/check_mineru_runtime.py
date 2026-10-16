"""Synthetic offline acceptance of the exact portable MinerU worker.

No private inputs, API credentials, network/model calls outside the local parser.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from pathlib import Path
import subprocess
import time
import uuid

ROOT = Path(__file__).resolve().parents[1]


def fixtures(directory: Path) -> list[tuple[str, Path]]:
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from pptx import Presentation
    from pptx.util import Inches
    directory.mkdir(parents=True, exist_ok=True)
    font_path = Path(os.environ["WINDIR"]) / "Fonts/segoeui.ttf"
    image = Image.new("RGB", (1500, 1800), "white")
    draw = ImageDraw.Draw(image)
    title = ImageFont.truetype(str(font_path), 56)
    body = ImageFont.truetype(str(font_path), 38)
    draw.text((100, 80), "DeepSteam OFFLINE 12345", font=title, fill="black")
    draw.text((100, 210), "Offline scan reading acceptance", font=title, fill="black")
    draw.text((100, 330), "This synthetic fixture contains no private learning material.", font=body, fill="black")
    draw.text((540, 470), "E = mc²", font=title, fill="black")
    draw.text((430, 610), "P(A) = 2 / 5 = 0.4", font=body, fill="black")
    table = [["Sample", "Count", "Score"], ["A", "12", "0.4"], ["B", "18", "0.6"]]
    for y in range(4):
        draw.line((120, 800 + y * 130, 1380, 800 + y * 130), fill="black", width=4)
    for x in range(4):
        draw.line((120 + x * 420, 800, 120 + x * 420, 1190), fill="black", width=4)
    for y, row in enumerate(table):
        for x, text in enumerate(row):
            draw.text((145 + x * 420, 830 + y * 130), text, font=body, fill="black")
    draw.rectangle((180, 1330, 580, 1570), fill=(204, 95, 50))
    draw.rectangle((680, 1330, 1080, 1570), fill=(40, 100, 190))
    draw.text((180, 1610), "Figure 1: orange and blue blocks", font=body, fill="black")
    png = directory / "synthetic-scan.png"
    image.save(png)
    pdf = directory / "synthetic-scan.pdf"
    image.save(pdf, resolution=150)
    doc = Document()
    doc.add_heading("DeepSteam DOCX OFFLINE 23456", 0)
    doc.add_paragraph("Document text: retain images, formulas, and tables.")
    omath = OxmlElement("m:oMath")
    def math_run(value):
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = value
        run.append(text)
        return run
    omath.append(math_run("E=m"))
    superscript = OxmlElement("m:sSup")
    base = OxmlElement("m:e")
    base.append(math_run("c"))
    exponent = OxmlElement("m:sup")
    exponent.append(math_run("2"))
    superscript.extend([base, exponent])
    omath.append(superscript)
    doc.add_paragraph()._p.append(omath)
    grid = doc.add_table(rows=3, cols=3)
    for y, row in enumerate(table):
        for x, text in enumerate(row):
            grid.cell(y, x).text = text
    doc.add_picture(str(png), width=__import__("docx").shared.Inches(3))
    doc_path = directory / "synthetic.docx"
    doc.save(doc_path)
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[5])
    slide.shapes.title.text = "DeepSteam PPTX OFFLINE 34567"
    box = slide.shapes.add_textbox(Inches(.4), Inches(1.0), Inches(8), Inches(.6))
    box.text_frame.text = "Slide text E=mc²"
    grid = slide.shapes.add_table(3, 3, Inches(.4), Inches(2.0), Inches(5), Inches(1.5)).table
    for y, row in enumerate(table):
        for x, text in enumerate(row):
            grid.cell(y, x).text = text
    slide.shapes.add_picture(str(png), Inches(6), Inches(1.5), height=Inches(4))
    ppt_path = directory / "synthetic.pptx"
    ppt.save(ppt_path)
    return [("scan_pdf", pdf), ("docx", doc_path), ("pptx", ppt_path)]


def run(tier: str, label: str) -> dict:
    base = ROOT / "runtime/mineru-acceptance" / label
    inputs = fixtures(base / "inputs")
    python = ROOT / "runtime/mineru/python.exe"
    worker = ROOT / "scripts/mineru_document_worker.py"
    worker_before = hashlib.sha256(worker.read_bytes()).hexdigest()
    results = []
    for kind, path in inputs:
        request = {"schema_version": 1, "action": "parse", "source_path": str(path),
                   "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                   "output_dir": str(base / (kind + "-output")), "request_id": label + "-" + kind,
                   "engine_options": {"tier": tier}}
        start = time.monotonic()
        environment = os.environ.copy()
        environment["PYTHONDONTWRITEBYTECODE"] = "1"
        result = subprocess.run([str(python), str(worker)], input=json.dumps(request),
                                encoding="utf-8", errors="replace", capture_output=True, env=environment,
                                timeout=1200)
        (base / (kind + ".stderr.txt")).write_text(result.stderr, encoding="utf-8")
        (base / (kind + ".stdout.json")).write_text(result.stdout, encoding="utf-8")
        try:
            parsed = json.loads(result.stdout)
        except ValueError:
            parsed = {"ok": False, "error": {"code": "invalid_worker_response", "message": result.stdout[-500:]}}
        output = Path(request["output_dir"])
        md = (output / "markdown.md").read_text("utf-8") if (output / "markdown.md").exists() else ""
        expected = {"scan_pdf": "12345", "docx": "23456", "pptx": "34567"}[kind]
        checks = {"worker_completed": result.returncode == 0 and parsed.get("ok") is True,
                  "expected_text": expected in md, "recognized_fixture_text": "Document text" in md,
                  "table_values": "12" in md and "18" in md and "0.4" in md,
                  "image_assets": parsed.get("coverage", {}).get("asset_count", 0) > 0,
                  "full_document": parsed.get("coverage", {}).get("full_document") is True,
                  "no_outbound_attempts": parsed.get("blocked_network_attempts") == 0,
                  "source_unchanged": hashlib.sha256(path.read_bytes()).hexdigest() == request["source_sha256"]}
        equation_text = " ".join(str(b.get("content", "")) for u in parsed.get("units", []) for b in u.get("equations", []))
        if kind in {"scan_pdf", "docx"}:
            compact_equation = re.sub(r"\s+", "", equation_text)
            checks["formula_superscript_retained"] = "^{2}" in compact_equation or "^2" in compact_equation
        if kind == "scan_pdf":
            checks["full_page_pixels_retained"] = parsed.get("coverage", {}).get("source_page_previews") == 1
        entry = {"kind": kind, "tier": tier, "elapsed_seconds": round(time.monotonic() - start, 3),
                 "checks": checks, "passed": all(checks.values()), "coverage": parsed.get("coverage"),
                 "formula_observation": [block for u in parsed.get("units", []) for block in u.get("equations", [])],
                 "quality_observations": (["Formula has additional absolute-value delimiters absent from the synthetic source; source pixels must be reviewed."]
                    if "\\left|" in equation_text or "\\right|" in equation_text else []),
                 "runtime_policy": parsed.get("runtime_policy"),
                 "error": parsed.get("error"), "output_dir": str(output)}
        results.append(entry)
        print(json.dumps(entry, ensure_ascii=False), flush=True)
    worker_after = hashlib.sha256(worker.read_bytes()).hexdigest()
    report = {"schema": "deepsteam-mineru-offline-acceptance-v1", "passed": all(x["passed"] for x in results) and worker_before == worker_after,
              "network_policy": "local model config plus outbound Python socket/DNS denial; no remote provider",
              "synthetic_only": True, "private_inputs_used": False, "tier": tier, "results": results,
              "acceptance_scope": "Offline execution, structure, source preservation and explicit quality limits; not a certification of perfect OCR.",
              "worker_sha256": worker_before, "worker_unchanged": worker_before == worker_after}
    (ROOT / ("reports/mineru-offline-acceptance-" + label + ".json")).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--tier", choices=["basic", "standard", "advanced"], default="basic")
    parser.add_argument("--label", default="run-" + uuid.uuid4().hex[:8])
    args = parser.parse_args()
    result = run(args.tier, args.label)
    raise SystemExit(0 if result["passed"] else 1)
