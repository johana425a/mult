"""Synthetic native ingest regressions. No account, provider, or user data."""
from __future__ import annotations

import asyncio
import base64
import hashlib
import io
import json
import os
import sys
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from integration import shared_document_extraction as shared

ROOT = Path(__file__).resolve().parents[1]
PORTABLE = Path(os.environ.get('DEEPSTEAM_TEST_PYTHON', sys.executable))


def row(text, index=1, **extras):
    return {'kind': 'page', 'text': text, 'locator': {'page': index, 'page_kind': 'pdf_physical'},
            'assets': [], 'char_start': 0, 'char_end': len(text), 'total_chars': len(text),
            'truncated_unit': False, **extras}


def response(rows, count=1, **extra):
    return {'ok': True, 'document': {'sha256': 'a' * 64, 'bytes': 20, 'unit_count': count},
            'units': rows, 'warnings': [], 'status': 'ready', 'truncated': False, 'next_cursor': None, **extra}


class CompletenessTests(unittest.TestCase):
    def test_bundled_readiness_reports_repair_instead_of_api_configuration(self):
        from deeptutor.services.parsing.engines.text_only.engine import TextOnlyParser

        with patch.object(shared, 'runtime_ready', return_value=True):
            self.assertTrue(TextOnlyParser().is_ready({}).ready)
        with patch.object(shared, 'runtime_ready', return_value=False):
            report = TextOnlyParser().is_ready({})
            self.assertFalse(report.ready)
            self.assertEqual(report.reason, 'local_reader_missing')
            self.assertIn('Repair', report.message)

    def test_native_advertisement_and_validation_include_images_without_legacy_word(self):
        from deeptutor.api.routers.reading import supported_formats
        from deeptutor.services.parsing.engines.text_only.engine import TextOnlyParser
        from deeptutor.services.rag.file_routing import FileTypeRouter
        from deeptutor.utils.document_validator import DocumentValidator

        advertised = asyncio.run(supported_formats())
        self.assertIn('.png', advertised.extensions)
        self.assertNotIn('.doc', advertised.extensions)
        self.assertIn('.png', TextOnlyParser().supported_formats())
        self.assertTrue(FileTypeRouter.needs_parser('fixture.png'))
        self.assertEqual(DocumentValidator.validate_upload_safety('fixture.png', 1000), 'fixture.png')

    def test_cursor_fragments_join_before_native_ingest(self):
        first = row('abc', char_end=3, total_chars=6, truncated_unit=True)
        second = row('def', char_start=3, char_end=6, total_chars=6)
        with patch.object(shared, '_request', side_effect=[response([first], truncated=True, next_cursor='next'), response([second])]):
            result = shared.extract_document(ROOT / 'runtime/fixture.pdf', include_assets=False)
        self.assertEqual(result['units'][0]['text'], 'abcdef')
        self.assertTrue(result['coverage']['document_fully_read'])

    def test_partial_batch_retries_ocr_and_preserves_later_pages(self):
        pending = row('', index=2, status='needs_ocr', assets=[{'id': 'pdf-page:2'}])
        with patch.object(shared, '_request', side_effect=[response([row('one'), pending], 2), response([row('recognized', index=2, extraction='ocr')], 2)]) as request:
            result = shared.extract_document(ROOT / 'runtime/fixture.pdf', include_assets=False)
        self.assertEqual([r['text'] for r in result['units']], ['one', 'recognized'])
        self.assertEqual(request.call_args.args[1]['start_unit'], 2)
        self.assertEqual(request.call_args.args[1]['document_sha256'], 'a' * 64)

    def test_missing_or_malformed_visual_evidence_still_fails(self):
        pending = response([row('', status='needs_ocr', assets=[{'id': 'pdf-page:1'}])])
        with patch.object(shared, '_request', return_value=pending), self.assertRaises(shared.SharedDocumentError):
            shared.extract_document(ROOT / 'runtime/fixture.pdf', include_assets=False)
        corrupt = response([], asset={'mime_type': 'image/png', 'base64': base64.b64encode(b'\x89PNG\r\n\x1a\nnot-a-png').decode()})
        with patch.object(shared, '_request', side_effect=[pending, pending, corrupt]), self.assertRaisesRegex(shared.SharedDocumentError, 'corrupt'):
            shared.extract_document(ROOT / 'runtime/fixture.pdf', include_assets=False)

    def test_textless_valid_image_is_retained_as_partial_even_without_optional_previews(self):
        from PIL import Image

        image = io.BytesIO()
        Image.new('RGB', (100, 100), 'blue').save(image, format='PNG')
        pending = response([row('', status='needs_ocr', assets=[{'id': 'pdf-page:1'}])])
        picture = response([], asset={'mime_type': 'image/png', 'base64': base64.b64encode(image.getvalue()).decode()})
        with patch.object(shared, '_request', side_effect=[pending, pending, picture]):
            result = shared.extract_document(ROOT / 'runtime/fixture.pdf', include_assets=False)
        self.assertEqual(result['status'], 'partial')
        self.assertEqual(result['units'][0]['extraction'], 'visual_only')
        self.assertTrue(result['units'][0]['text_is_placeholder'])
        self.assertEqual(result['images'][0]['data'], image.getvalue())
        self.assertIn('text_not_recognized', result['coverage']['unresolved'])

    def test_missing_bundled_components_fail_before_visual_fallback(self):
        with patch.object(shared, 'runtime_ready', return_value=False), self.assertRaisesRegex(shared.SharedDocumentError, 'Repair the installation'):
            shared._request(ROOT / 'runtime/fixture.pdf', {'action': 'read'}, 5)

    def test_limits_missing_pages_source_change_and_repeated_cursor_fail(self):
        cases = [
            [response([row('one')], count=2)],
            [response([row('one')], count=2, truncated=True, next_cursor='n'), response([row('two', index=2)], count=2, document={'sha256': 'b' * 64, 'unit_count': 2})],
            [response([row('a')], count=2, truncated=True, next_cursor='n'), response([row('b', index=2)], count=2, truncated=True, next_cursor='n')],
        ]
        for values in cases:
            with self.subTest(values=values), patch.object(shared, '_request', side_effect=values), self.assertRaises(shared.SharedDocumentError):
                shared.extract_document(ROOT / 'runtime/fixture.pdf', include_assets=False)
        with patch.object(shared, 'MAX_CHARS', 2), patch.object(shared, '_request', return_value=response([row('long')])), self.assertRaises(shared.SharedDocumentError):
            shared.extract_document(ROOT / 'runtime/fixture.pdf', include_assets=False)

    def test_fidelity_metadata_survives_native_parser_artifact(self):
        equation = {'format': 'latex', 'text': r'\frac{x}{y}', 'status': 'complete'}
        located = row('formula', equations=[equation], table={'rows': [['a', 'b']]}, coverage={'unresolved': ['unknown_layout']})
        located['locator'] = {'block': 1, 'paragraph': 1}
        with tempfile.TemporaryDirectory(dir=ROOT / 'runtime') as directory, patch.object(shared, '_request', return_value=response([located], warnings=['Review source image'])):
            result = shared.extract_document(Path(directory) / 'fixture.docx', include_assets=False)
            shared.write_parser_artifacts(result, Path(directory) / 'parsed', 'fixture')
            blocks = json.loads((Path(directory) / 'parsed/fixture_content_list.json').read_text(encoding='utf-8'))
            self.assertEqual(blocks[0]['equations'], [equation])
            self.assertNotIn('page_idx', blocks[0])
            self.assertEqual(blocks[0]['table']['rows'], [['a', 'b']])
            self.assertEqual(result['status'], 'partial')
            self.assertIn('Review source image', (Path(directory) / 'parsed/fixture.md').read_text(encoding='utf-8'))


@unittest.skipUnless(PORTABLE.is_file(), 'packaged OCR Python test runtime unavailable')
class NativePipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='native-reader-test-', dir=ROOT / 'runtime')
        self.directory = Path(self.temp.name)
        self.runtime = patch.object(shared, 'runtime_paths', return_value=(PORTABLE, ROOT / 'scripts/workspace_document_reader.py'))
        self.runtime.start()

    def tearDown(self):
        self.runtime.stop()
        self.temp.cleanup()

    def test_packaged_runtime_is_ready_without_any_provider(self):
        from deeptutor.services.parsing.engines.text_only.engine import TextOnlyParser

        self.assertTrue(shared.runtime_ready())
        self.assertTrue(TextOnlyParser().is_ready({}).ready)

    def test_diagram_only_pdf_retains_partial_visual_in_native_reading_and_parser(self):
        import pymupdf
        from deeptutor.reading.extract import extract_material
        from deeptutor.services.parsing.engines.text_only.engine import TextOnlyParser

        path = self.directory / 'diagram.pdf'
        document = pymupdf.open()
        page = document.new_page(width=400, height=400)
        page.draw_circle((130, 130), 75, color=(0, 0, 1), fill=(0, 0, 1))
        page.draw_rect(pymupdf.Rect(210, 210, 330, 330), color=(1, 0, 0), fill=(1, 0, 0))
        document.save(path)
        document.close()
        reading = extract_material(path)
        self.assertIn('no text was recognized', reading.units[0])
        self.assertTrue(reading.media)
        self.assertEqual(reading.media[-1].locator, 1)
        output = self.directory / 'diagram-parser'
        TextOnlyParser().parse(path, output, config={})
        record = json.loads((output / 'source-extraction.json').read_text(encoding='utf-8'))
        self.assertEqual(record['status'], 'partial')
        self.assertIn('text_not_recognized', record['coverage']['unresolved'])
        self.assertTrue(list((output / 'images').glob('*.png')))

    def scan_pdf(self):
        from PIL import Image, ImageDraw, ImageFont
        import pymupdf

        image = Image.new('RGB', (1100, 400), 'white')
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', 50)
        draw.text((40, 80), 'LOCAL SCAN EVIDENCE 12345', fill='black', font=font)
        stream = io.BytesIO()
        image.save(stream, format='PNG')
        (self.directory / 'scan.png').write_bytes(stream.getvalue())
        document = pymupdf.open()
        page = document.new_page(width=700, height=300)
        page.insert_image(page.rect, stream=stream.getvalue())
        path = self.directory / 'scan.pdf'
        document.save(path)
        document.close()
        return path

    def test_standalone_scan_image_uses_same_local_ocr(self):
        from deeptutor.reading.extract import extract_material
        from deeptutor.services.parsing.engines.text_only.engine import TextOnlyParser

        self.scan_pdf()
        image = self.directory / 'scan.png'
        reading = extract_material(image)
        self.assertIn('12345', '\n'.join(reading.units))
        self.assertTrue(reading.media)
        output = self.directory / 'image-parser'
        TextOnlyParser().parse(image, output, config={})
        self.assertIn('12345', (output / 'scan.md').read_text(encoding='utf-8'))

    def test_scan_reaches_native_reading_and_parser_without_provider(self):
        from deeptutor.reading.extract import extract_material
        from deeptutor.services.parsing.engines.text_only.engine import TextOnlyParser

        path = self.scan_pdf()
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        extraction = extract_material(path)
        self.assertEqual(extraction.unit, 'page')
        self.assertEqual(extraction.render_mode, 'pdf')
        self.assertTrue(extraction.has_raw_view)
        self.assertIn('12345', extraction.units[0])
        output = self.directory / 'parser'
        TextOnlyParser().parse(path, output, config={})
        self.assertIn('12345', (output / 'scan.md').read_text(encoding='utf-8'))
        blocks = json.loads((output / 'scan_content_list.json').read_text(encoding='utf-8'))
        self.assertEqual(blocks[0]['source_locator']['page'], 1)
        self.assertEqual(blocks[0]['extraction'], 'ocr')
        self.assertTrue(list((output / 'images').glob('*.png')))
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), digest)

    def test_docx_formula_table_and_image_survive_both_native_pipelines(self):
        from docx import Document
        from docx.oxml import parse_xml
        from PIL import Image
        from deeptutor.reading.extract import extract_material
        from deeptutor.services.parsing.engines.text_only.engine import TextOnlyParser

        document = Document()
        p = document.add_paragraph('Synthetic equation: ')
        p._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>x</m:t></m:r></m:num><m:den><m:r><m:t>y</m:t></m:r></m:den></m:f></m:oMath>'))
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'table left'
        table.cell(0, 1).text = 'table right'
        stream = io.BytesIO()
        Image.new('RGB', (150, 120), 'blue').save(stream, format='PNG')
        p = document.add_paragraph('Illustration alongside known text')
        p.add_run().add_picture(io.BytesIO(stream.getvalue()))
        path = self.directory / 'fixture.docx'
        document.save(path)
        reading = extract_material(path)
        self.assertIn(r'\frac{x}{y}', '\n'.join(reading.units))
        self.assertIn('table left', '\n'.join(reading.units))
        self.assertTrue(reading.media)
        for picture in reading.media:
            self.assertIn(picture.name, reading.units[picture.locator - 1])
        output = self.directory / 'parser'
        TextOnlyParser().parse(path, output, config={})
        blocks = json.loads((output / 'fixture_content_list.json').read_text(encoding='utf-8'))
        self.assertTrue(any(b.get('equations') for b in blocks))
        self.assertTrue(any(b.get('table') for b in blocks))
        self.assertTrue(any(b['type'] == 'image' for b in blocks))


if __name__ == '__main__':
    unittest.main()
