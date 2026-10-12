"""Synthetic-only contract and security tests for the shared workspace reader."""
from __future__ import annotations

import base64
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock
import zipfile

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts/workspace_document_reader.py"
spec = importlib.util.spec_from_file_location("workspace_document_reader", SCRIPT)
reader = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(reader)


def png_bytes() -> bytes:
    from PIL import Image, ImageDraw
    image = Image.new("RGB", (320, 160), "white")
    ImageDraw.Draw(image).text((30, 50), "Local document fixture", fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


class ReaderTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="workspace-reader-", dir=ROOT / "runtime")
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def read(self, path, **kwargs):
        return reader.handle_request({"root": str(self.root), "path": path, "ocr": "off", **kwargs})

    def pdf(self, name="text.pdf", *, scan=False, blank=False):
        import pymupdf
        doc = pymupdf.open()
        for n in range(3):
            page = doc.new_page()
            if scan:
                page.insert_image(pymupdf.Rect(30, 30, 350, 190), stream=png_bytes())
            elif not blank:
                page.insert_text((30, 50), f"Physical page {n + 1} synthetic text")
        path = self.root / name
        doc.save(path)
        doc.close()
        return path

    def docx(self):
        from docx import Document
        document = Document()
        document.add_paragraph("First paragraph synthetic paragraph")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "table A"
        table.cell(0, 1).text = "table B"
        document.add_picture(io.BytesIO(png_bytes()))
        document.add_paragraph("Last paragraph")
        document.save(self.root / "word.docx")

    def pptx(self):
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = "Slide two"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        table.cell(0, 0).text = "table left"
        table.cell(0, 1).text = "table right"
        slide.notes_slide.notes_text_frame.text = "Teacher notes"
        slide.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(3))
        deck.save(self.root / "slides.pptx")

    def test_pdf_pages_and_cursor_are_physical(self):
        self.pdf()
        first = self.read("text.pdf", max_units=1)
        self.assertTrue(first["ok"])
        self.assertTrue(first["truncated"])
        self.assertEqual(first["document"]["unit_count"], 3)
        self.assertEqual(first["units"][0]["locator"], {"page": 1, "page_kind": "pdf_physical"})
        second = self.read("text.pdf", cursor=first["next_cursor"], max_units=1)
        self.assertIn("Physical page 2", second["text"])
        self.assertEqual(second["units"][0]["locator"]["page"], 2)

    def test_pdf_does_not_extract_unrequested_pages(self):
        self.pdf()
        original = reader.Document.get
        visited = []
        def observed(document, index):
            visited.append(index)
            return original(document, index)
        with mock.patch.object(reader.Document, "get", observed):
            result = self.read("text.pdf", max_units=1, start_unit=3)
        self.assertEqual(visited, [2])
        self.assertFalse(result["truncated"])

    def test_pdf_asset_requires_version_and_is_png(self):
        self.pdf()
        first = self.read("text.pdf")
        bad = self.read("text.pdf", action="asset", asset_id="pdf-page:1")
        self.assertEqual(bad["error"]["code"], "SOURCE_VERSION_REQUIRED")
        result = self.read("text.pdf", action="asset", asset_id="pdf-page:1", document_sha256=first["document_sha256"], max_dimension=512)
        self.assertTrue(result["ok"])
        self.assertEqual(result["asset"]["mime_type"], "image/png")
        self.assertLessEqual(max(result["asset"]["width"], result["asset"]["height"]), 512)
        self.assertTrue(base64.b64decode(result["asset"]["base64"]).startswith(b"\x89PNG"))

    def test_scan_not_falsely_claimed_read(self):
        self.pdf("scan.pdf", scan=True)
        result = self.read("scan.pdf", max_units=1)
        self.assertEqual(result["status"], "needs_ocr")
        self.assertEqual(result["text"], "")
        self.assertEqual(result["units"][0]["extraction"], "none")
        self.assertEqual(result["units"][0]["assets"][0]["id"], "pdf-page:1")

    def test_blank_pdf_is_empty_not_ocr_or_ready(self):
        self.pdf("blank.pdf", blank=True)
        result = self.read("blank.pdf")
        self.assertEqual(result["status"], "empty")

    def test_ocr_result_label_is_explicit(self):
        self.pdf("scan.pdf", scan=True)
        with mock.patch.object(reader.Document, "local_ocr", return_value="Recognized synthetic text"):
            result = self.read("scan.pdf", max_units=1, ocr="auto")
        self.assertEqual(result["units"][0]["extraction"], "ocr")
        self.assertTrue(any("OCR" in warning for warning in result["warnings"]))

    def test_ocr_unavailable_is_not_empty_success(self):
        self.pdf("scan.pdf", scan=True)
        with mock.patch.dict(sys.modules, {"rapidocr_onnxruntime": None, "rapidocr": None}):
            result = self.read("scan.pdf", max_units=1, ocr="auto")
        self.assertEqual(result["status"], "needs_ocr")
        self.assertTrue(any("unavailable" in warning for warning in result["warnings"]))

    def test_ocr_page_budget_does_not_start_a_third_page(self):
        path = self.pdf("scan.pdf", scan=True)
        document = reader.Document(path.read_bytes(), ".pdf", "auto")
        document.ocr_attempts = 2
        try:
            with mock.patch.dict(sys.modules, {"rapidocr": None, "rapidocr_onnxruntime": None}):
                result = document.get(2)
            self.assertEqual(result["status"], "needs_ocr")
            self.assertEqual(document.ocr_attempts, 2)
            self.assertTrue(any("limited to two pages" in warning for warning in document.warnings))
        finally:
            document.close()

    def test_word_paragraph_table_and_image_order(self):
        self.docx()
        result = self.read("word.docx")
        self.assertTrue(result["ok"], result)
        self.assertIn("First paragraph", result["text"])
        self.assertIn("table A\ttable B", result["text"])
        self.assertEqual([row["kind"] for row in result["units"]], ["paragraph", "table", "paragraph", "paragraph"])
        image_row = result["units"][2]
        self.assertEqual(image_row["locator"], {"paragraph": 2})
        result = self.read("word.docx", action="asset", asset_id=image_row["assets"][0]["id"], document_sha256=result["document_sha256"])
        self.assertTrue(result["ok"], result)

    def test_ppt_blank_slide_preserves_number_notes_table_images(self):
        self.pptx()
        result = self.read("slides.pptx")
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["document"]["unit_count"], 2)
        self.assertEqual(result["units"][0]["locator"], {"slide": 1})
        second = result["units"][1]
        self.assertEqual(second["locator"], {"slide": 2})
        self.assertIn("Teacher notes", second["text"])
        self.assertIn("table left\ttable right", second["text"])
        self.assertEqual(len(second["assets"]), 1)

    def test_xlsx_formula_not_executed_and_sheet_source(self):
        from openpyxl import Workbook
        doc = Workbook()
        sheet = doc.active
        sheet.title = "Research"
        sheet.append(["A", "=1+1"])
        doc.save(self.root / "table.xlsx")
        result = self.read("table.xlsx")
        self.assertEqual(result["units"][0]["locator"], {"sheet": "Research", "row": 1})
        self.assertIn("=1+1", result["text"])

    def test_raw_image_has_bounded_png_asset(self):
        (self.root / "figure.png").write_bytes(png_bytes())
        result = self.read("figure.png")
        self.assertEqual(result["units"][0]["assets"][0]["id"], "image:1")
        self.assertEqual(result["units"][0]["extraction"], "none")
        output = self.read("figure.png", action="asset", asset_id="image:1", document_sha256=result["document_sha256"])
        self.assertTrue(output["ok"], output)
        self.assertTrue(base64.b64decode(output["asset"]["base64"]).startswith(b"\x89PNG"))

    def test_markdown_relative_images_wiki_and_changed_image(self):
        folder = self.root / "book"
        folder.mkdir()
        (self.root / "figure.png").write_bytes(png_bytes())
        (folder / "page.md").write_text("![](../figure.png)\n![[../figure.png|400]]\n", encoding="utf-8")
        result = self.read("book/page.md")
        self.assertEqual(len(result["units"][0]["assets"]), 1)
        self.assertEqual(len(result["units"][1]["assets"]), 1)
        asset_id = result["units"][0]["assets"][0]["id"]
        output = self.read("book/page.md", action="asset", asset_id=asset_id, document_sha256=result["document_sha256"])
        self.assertTrue(output["ok"], output)
        from PIL import Image
        Image.new("RGB", (10, 10), "red").save(self.root / "figure.png")
        changed = self.read("book/page.md", action="asset", asset_id=asset_id, document_sha256=result["document_sha256"])
        self.assertEqual(changed["error"]["code"], "SOURCE_CHANGED")

    def test_markdown_does_not_fetch_external_or_escaping_images(self):
        (self.root / "page.md").write_text("![](https://example.invalid/private.png)\n![](../private.png)\n", encoding="utf-8")
        with mock.patch("urllib.request.urlopen", side_effect=AssertionError("must never fetch")):
            result = self.read("page.md")
        self.assertTrue(result["ok"], result)
        self.assertTrue(all(not unit["assets"] for unit in result["units"]))
        self.assertTrue(any("Remote image" in warning for warning in result["warnings"]))
        self.assertTrue(any("workspace" in warning for warning in result["warnings"]))

    def test_markdown_reads_only_current_window_images(self):
        (self.root / "page.md").write_text("Text first\n![](missing.png)\n", encoding="utf-8")
        result = self.read("page.md", max_units=1)
        self.assertTrue(result["truncated"])
        self.assertEqual(result["warnings"], [])

    def test_markdown_long_line_pagination_is_lossless(self):
        source = "Title\n" + "X" * 800 + "\nLast line\n"
        (self.root / "notes.md").write_bytes(source.encode("utf-8"))
        cursor = None
        chunks = []
        for _ in range(10):
            result = self.read("notes.md", max_chars=256, cursor=cursor)
            self.assertTrue(result["ok"], result)
            chunks.extend(row["text"] for row in result["units"])
            cursor = result["next_cursor"]
            if not cursor:
                break
        self.assertEqual("".join(chunks), source)

    def test_utf8_utf16_and_gbk(self):
        for encoding in ["utf-8-sig", "utf-16", "gb18030"]:
            (self.root / "text.txt").write_bytes("Local text\nSecond line".encode(encoding))
            self.assertIn("Local text", self.read("text.txt")["text"])

    def test_changed_source_rejects_cursor_and_asset_version(self):
        (self.root / "notes.txt").write_text("A\nB\n", encoding="utf-8")
        first = self.read("notes.txt", max_units=1)
        (self.root / "notes.txt").write_text("Changed\n", encoding="utf-8")
        self.assertEqual(self.read("notes.txt", cursor=first["next_cursor"])["error"]["code"], "SOURCE_CHANGED")
        self.assertEqual(self.read("notes.txt", action="asset", asset_id="anything", document_sha256=first["document_sha256"])["error"]["code"], "SOURCE_CHANGED")

    def test_paths_cannot_escape_or_use_ads(self):
        for path in ["../secret.txt", "..\\secret.txt", "D:\\secret.txt", "/secret.txt", "\\\\server\\share\\secret.txt", "notes.txt:stream", "nested/../notes.txt"]:
            result = self.read(path)
            self.assertFalse(result["ok"], path)
            self.assertEqual(result["error"]["code"], "PATH_OUTSIDE_WORKSPACE")
            self.assertNotIn(str(self.root), json.dumps(result))

    def test_missing_path_error_does_not_leak_absolute_path(self):
        result = self.read("private-note.txt")
        self.assertEqual(result["error"]["code"], "FILE_NOT_FOUND")
        self.assertNotIn("private-note", json.dumps(result))

    def test_symlink_rejected(self):
        original = self.root / "source.txt"
        original.write_text("source")
        try:
            (self.root / "link.txt").symlink_to(original)
        except (OSError, NotImplementedError):
            self.skipTest("OS does not permit fixture symlinks")
        self.assertEqual(self.read("link.txt")["error"]["code"], "PATH_REPARSE_POINT")

    def test_simulated_windows_reparse_rejected(self):
        class JunctionStat:
            st_mode = stat_mode = 0o040755
            st_file_attributes = 0x400
        with mock.patch.object(Path, "lstat", return_value=JunctionStat()):
            with self.assertRaises(reader.ReaderError) as raised:
                reader._plain_ancestors(self.root)
        self.assertEqual(raised.exception.code, "PATH_REPARSE_POINT")

    def test_legacy_office_has_clear_conversion_error(self):
        for name in ["old.doc", "old.ppt", "old.xls"]:
            (self.root / name).write_bytes(b"\xd0\xcf\x11\xe0legacy")
            self.assertEqual(self.read(name)["error"]["code"], "CONVERSION_REQUIRED")

    def test_format_spoof_and_binary_text_rejected(self):
        for name, content in [("wrong.pdf", b"not PDF"), ("wrong.docx", b"not office"), ("wrong.md", b"%PDF-1.7\x00"), ("binary.txt", b"A\x00B")]:
            (self.root / name).write_bytes(content)
            self.assertEqual(self.read(name)["error"]["code"], "FORMAT_MISMATCH")

    def archive(self, entries):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as doc:
            for name, value in entries:
                doc.writestr(name, value)
        (self.root / "bad.docx").write_bytes(stream.getvalue())

    def test_archive_traversal_and_entity_rejected(self):
        for entries in [[("../bad", "x")], [("word/document.xml", '<!DOCTYPE x [<!ENTITY xx SYSTEM "file:///secret">]>')]]:
            self.archive(entries)
            self.assertEqual(self.read("bad.docx")["error"]["code"], "UNSAFE_ARCHIVE")

    def test_archive_compression_bomb_rejected(self):
        self.archive([("word/document.xml", "X" * (2 * 1024 * 1024))])
        self.assertEqual(self.read("bad.docx")["error"]["code"], "ARCHIVE_LIMIT")

    def test_utf16_xml_external_entity_rejected(self):
        self.archive([("word/document.xml", '<!DOCTYPE x [<!ENTITY xx SYSTEM "file:///secret">]>'.encode("utf-16"))])
        self.assertEqual(self.read("bad.docx")["error"]["code"], "UNSAFE_ARCHIVE")

    def test_file_size_budget(self):
        (self.root / "small.txt").write_bytes(b"x" * 257)
        with mock.patch.object(reader, "MAX_TEXT_BYTES", 256):
            self.assertEqual(self.read("small.txt")["error"]["code"], "FILE_TOO_LARGE")

    def test_asset_id_is_not_file_path(self):
        self.docx()
        result = self.read("word.docx")
        asset = self.read("word.docx", action="asset", asset_id="../../private.png", document_sha256=result["document_sha256"])
        self.assertEqual(asset["error"]["code"], "ASSET_NOT_FOUND")

    def test_empty_text_and_invalid_options(self):
        (self.root / "empty.md").write_text("")
        self.assertEqual(self.read("empty.md")["status"], "empty")
        for options in [{"max_units": 41}, {"max_chars": 200000}, {"ocr": "cloud"}, {"action": "execute"}]:
            self.assertFalse(self.read("empty.md", **options)["ok"])

    def test_cli_is_one_json_and_original_is_unchanged(self):
        path = self.root / "source.md"
        path.write_text("Synthetic test document only\n", encoding="utf-8")
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        request = {"root": str(self.root), "path": path.name, "ocr": "off"}
        result = subprocess.run([sys.executable, "-I", str(SCRIPT)], input=json.dumps(request).encode(), capture_output=True, timeout=20)
        self.assertEqual(result.returncode, 0)
        self.assertEqual(result.stderr, b"")
        output = json.loads(result.stdout)
        self.assertTrue(output["ok"], output)
        self.assertEqual(output["document_sha256"], before)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)


if __name__ == "__main__":
    unittest.main()
