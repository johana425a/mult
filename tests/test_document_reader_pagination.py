"""Pagination regressions with real PDF/DOCX files and synthetic OCR output.

The actual Document.local_ocr budget and PDF rasterization run unchanged. Only
the RapidOCR backend is replaced; no model weights or remote service run.
"""
from __future__ import annotations

import hashlib
import io
from pathlib import Path
import sys
import tempfile
from types import ModuleType, SimpleNamespace
import unittest
from unittest import mock

from scripts import workspace_document_reader as reader


ROOT = Path(__file__).resolve().parents[1]


class ReaderPaginationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(
            prefix='document-pagination-tests-', dir=ROOT / 'runtime')
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def read(self, path, **options):
        value = reader.handle_request({'root': str(self.root), 'path': path.name, **options})
        self.assertTrue(value['ok'], value)
        return value

    def test_twelve_scanned_pdf_pages_survive_two_page_ocr_budget_without_cursor_gaps(self):
        import pymupdf
        from PIL import Image, ImageDraw

        path = self.root / 'twelve-scanned-pages.pdf'
        document = pymupdf.open()
        for page_number in range(1, 13):
            # The unique corner colour lets the synthetic recognizer identify
            # the physical page from its actual rendered pixels, not call order.
            pixels = Image.new('RGB', (240, 120), (20 + page_number, 150, 200))
            ImageDraw.Draw(pixels).text((20, 45), f'Synthetic scan {page_number}', fill='black')
            stream = io.BytesIO()
            pixels.save(stream, format='PNG')
            page = document.new_page(width=240, height=120)
            page.insert_image(page.rect, stream=stream.getvalue())
            self.assertEqual(page.get_text().strip(), '')
        document.save(path)
        document.close()
        source_hash = hashlib.sha256(path.read_bytes()).hexdigest()

        package = self.root / 'synthetic_ocr_package'
        models = package / 'models'
        models.mkdir(parents=True)
        for filename in ('PP-OCRv6_det_small.onnx', 'PP-OCRv6_rec_small.onnx',
                         'ch_ppocr_mobile_v2.0_cls_mobile.onnx'):
            (models / filename).write_bytes(b'SYNTHETIC PLACEHOLDER; NOT A MODEL')
        batches = []

        class SyntheticRapidOCR:
            def __init__(self, *, params):
                self.pages = []
                batches.append(self.pages)

            def __call__(self, png):
                with Image.open(io.BytesIO(png)) as image:
                    page_number = image.convert('RGB').getpixel((0, 0))[0] - 20
                self.pages.append(page_number)
                return SimpleNamespace(txts=[f'SCANNED PAGE {page_number:02d}'])

        backend = ModuleType('rapidocr')
        backend.__file__ = str(package / '__init__.py')
        backend.RapidOCR = SyntheticRapidOCR
        rows, cursors, cursor = [], set(), None
        with mock.patch.dict(sys.modules, {'rapidocr': backend}):
            for request_number in range(6):
                options = {'ocr': 'auto', 'max_units': 10, 'max_chars': 12000}
                if cursor:
                    options['cursor'] = cursor
                result = self.read(path, **options)
                self.assertEqual(result['document']['unit_count'], 12)
                self.assertEqual(result['document_sha256'], source_hash)
                self.assertEqual(len(result['units']), 2)
                self.assertEqual(result['coverage']['ocr_units'], 2)
                for row in result['units']:
                    number = row['locator']['page']
                    expected = f'SCANNED PAGE {number:02d}'
                    self.assertEqual(row['locator']['page_kind'], 'pdf_physical')
                    self.assertEqual(row['unit_index'], number)
                    self.assertEqual(row['extraction'], 'ocr')
                    self.assertTrue(row['ocr_applied'])
                    self.assertEqual(row['text'], expected)
                    self.assertEqual(row['unit_text_sha256'], hashlib.sha256(expected.encode()).hexdigest())
                    self.assertEqual(row['assets'][0]['id'], f'pdf-page:{number}')
                rows.extend(result['units'])
                cursor = result['next_cursor']
                if request_number < 5:
                    self.assertTrue(result['truncated'])
                    self.assertIsNotNone(cursor)
                    self.assertNotIn(cursor, cursors)
                    cursors.add(cursor)
                    # The next unread physical page must be retried with a
                    # fresh OCR budget, not advanced past as an empty page.
                    self.assertEqual(reader.cursor_decode(cursor, source_hash),
                                     (2 * (request_number + 1), 0))
                else:
                    self.assertFalse(result['truncated'])
                    self.assertIsNone(cursor)
        self.assertEqual([row['locator']['page'] for row in rows], list(range(1, 13)))
        self.assertEqual(batches, [[1, 2], [3, 4], [5, 6], [7, 8], [9, 10], [11, 12]])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), source_hash)

    def test_long_paragraph_fragments_reassemble_exactly_and_keep_whole_unit_hash(self):
        from docx import Document

        text = ''.join(f'fragment-{index:04d} letters β 😀; ' for index in range(220))
        tail = 'A second paragraph must also be reached.'
        path = self.root / 'long-paragraph.docx'
        document = Document()
        document.add_paragraph(text)
        document.add_paragraph(tail)
        document.save(path)
        source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        expected_units = {1: text, 2: tail}
        fragments = {1: [], 2: []}
        ends = {1: 0, 2: 0}
        hashes = {1: set(), 2: set()}
        cursor, positions, requests = None, [], 0
        for _ in range(100):
            options = {'ocr': 'off', 'max_chars': 257, 'max_units': 10}
            if cursor:
                options['cursor'] = cursor
            result = self.read(path, **options)
            requests += 1
            self.assertEqual(result['document']['unit_count'], 2)
            self.assertEqual(result['document_sha256'], source_hash)
            self.assertLessEqual(sum(len(row['text']) for row in result['units']), 257)
            for row in result['units']:
                index = row['unit_index']
                self.assertEqual(row['locator'], {'paragraph': index})
                self.assertEqual(row['char_start'], ends[index])
                self.assertEqual(row['total_chars'], len(expected_units[index]))
                self.assertEqual(row['text'], expected_units[index][row['char_start']:row['char_end']])
                fragments[index].append(row['text'])
                ends[index] = row['char_end']
                hashes[index].add(row['unit_text_sha256'])
            cursor = result['next_cursor']
            if cursor is None:
                self.assertFalse(result['truncated'])
                break
            self.assertTrue(result['truncated'])
            position = reader.cursor_decode(cursor, source_hash)
            if positions:
                self.assertGreater(position, positions[-1])
            positions.append(position)
        else:
            self.fail('The paragraph cursor did not terminate')
        self.assertGreater(requests, 10)
        for index, expected in expected_units.items():
            self.assertEqual(''.join(fragments[index]), expected)
            self.assertEqual(ends[index], len(expected))
            self.assertEqual(hashes[index], {hashlib.sha256(expected.encode('utf-8')).hexdigest()})
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), source_hash)


if __name__ == '__main__':
    unittest.main()
