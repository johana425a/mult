"""Parse lane, background preparation and read-ahead with a deterministic
stand-in for the local MinerU worker.

The stand-in speaks the real ``--serve`` line protocol and writes real batch
results, but it runs no model: its "parsing" is a fixed sleep per page. The
timings below therefore measure scheduling (who waits for whom), never MinerU
or VLM speed.
"""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

from integration import document_jobs as module

PYTHON = Path(os.environ.get('DEEPSTEAM_TEST_PYTHON', sys.executable))

FIXTURE_WORKER = r'''
import json, os, pathlib, sys, time
import pymupdf
arguments = sys.argv[1:]
home = pathlib.Path(arguments[arguments.index('--home') + 1])
home.mkdir(parents=True, exist_ok=True)
settings = json.loads((home / 'fixture.json').read_text()) if (home / 'fixture.json').is_file() else {}
with (home / 'starts.txt').open('a') as log:
    log.write(json.dumps({'pid': os.getpid(), 'arguments': arguments}) + '\n')
time.sleep(settings.get('load_seconds', 0))
protocol = os.fdopen(os.dup(1), 'w', encoding='utf-8', buffering=1)
os.dup2(2, 1)
for line in sys.stdin:
    request = json.loads(line)
    identity = request['request_id']
    print(json.dumps({'event': 'progress', 'stage': 'parsing', 'request_id': identity}), file=sys.stderr, flush=True)
    source, output = pathlib.Path(request['source_path']), pathlib.Path(request['output_dir'])
    if settings.get('crash'):
        os._exit(3)
    if source.suffix.lower() == '.pdf':
        with pymupdf.open(source) as pdf:
            texts = [page.get_text().strip() for page in pdf]
    else:
        texts = ['IMAGE CONTENT']
    time.sleep(settings.get('page_seconds', 0) * len(texts))
    units = [{'text': 'PARSED ' + text, 'locator': {'kind': 'pdf_page', 'page': index + 1, 'page_idx': index}, 'assets': []}
             for index, text in enumerate(texts)]
    result = {'schema_version': 1, 'ok': True, 'status': 'ready', 'engine': 'mineru', 'source_sha256': request['source_sha256'],
              'units': units, 'warnings': [], 'coverage': {'full_document': True},
              'inference': {'mode': 'deterministic fixture', 'vlm_used': False}}
    (output / 'result.json').write_text(json.dumps(result))
    with (home / 'requests.jsonl').open('a') as log:
        log.write(json.dumps({'tier': request['engine_options']['tier'], 'pages': len(texts), 'at': time.time()}) + '\n')
    protocol.write(json.dumps({'ok': True, 'status': 'ready', 'engine': 'mineru', 'source_sha256': request['source_sha256'],
                               'result_path': 'result.json', 'request_id': identity}) + '\n')
'''


class ParseLaneTests(unittest.TestCase):
    def setUp(self):
        import pymupdf
        (module.ROOT / 'runtime').mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(prefix='parse-lane-', dir=module.ROOT / 'runtime')
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.workspace = self.root / 'workspace'
        self.workspace.mkdir()
        self.helper = self.root / 'fixture_mineru_worker.py'
        self.helper.write_text(FIXTURE_WORKER, encoding='utf-8')
        self.home = self.root / 'cache/document-jobs/mineru-worker'
        self.home.mkdir(parents=True)
        self.settings({})
        for name, count in (('long.pdf', 17), ('other.pdf', 9)):
            with pymupdf.open() as pdf:
                for number in range(1, count + 1):
                    pdf.new_page().insert_text((50, 50), f'{name} PAGE {number}')
                pdf.save(self.workspace / name)
        (self.workspace / 'notes.md').write_text('# Notes\nPlain text that never needs MinerU.', encoding='utf-8')
        self.threads = []
        self.lane_errors = []
        patches = [patch.object(module, 'mineru_runtime', return_value=(PYTHON, self.helper)),
                   patch.object(module.DocumentJobs, '_spawn', autospec=True, side_effect=self.spawn)]
        for item in patches:
            item.start()
            self.addCleanup(item.stop)
        self.addCleanup(self.join_lanes)
        self.jobs = self.broker()

    def broker(self):
        return module.DocumentJobs(test_root=self.root, python=PYTHON)

    def settings(self, value):
        (self.home / 'fixture.json').write_text(json.dumps(value), encoding='utf-8')

    def spawn(self, _jobs, _lock, flag):
        """In-process stand-in for a lane process, so the patches stay active."""
        def run():
            try:
                broker = self.broker()
                broker.drain_parses(idle_seconds=.3) if flag == '--drain-parses' else broker.drain()
            except BaseException as exc:  # Surface lane failures in the test.
                self.lane_errors.append(exc)
        thread = threading.Thread(target=run, name=f'synthetic-lane{flag}', daemon=True)
        self.threads.append(thread)
        thread.start()

    def join_lanes(self):
        for thread in list(self.threads):
            thread.join(30)
        self.assertEqual(self.lane_errors, [])

    def read(self, path, wait=0, **options):
        return self.jobs.wait(self.workspace, {'path': path, **options}, seconds=wait)

    def poll(self, path, job, condition, seconds=30):
        deadline = time.monotonic() + seconds
        while time.monotonic() < deadline:
            value = self.jobs.status(self.workspace, {'path': path, 'job_id': job['job_id']})
            if condition(value):
                return value
            time.sleep(.05)
        self.fail(f'Condition not reached for {path}: {value}')

    def parses(self):
        with self.jobs.connect() as db:
            return [dict(row) for row in db.execute('SELECT * FROM parses ORDER BY created,batch_start')]

    def starts(self):
        path = self.home / 'starts.txt'
        return path.read_text().splitlines() if path.is_file() else []

    def requests(self):
        path = self.home / 'requests.jsonl'
        return [json.loads(line) for line in path.read_text().splitlines()] if path.is_file() else []

    def test_other_files_are_read_while_a_batch_parses(self):
        self.settings({'page_seconds': .4})  # 4-page advanced batch: about 1.6 s of fixture "parsing"
        pending = self.read('long.pdf', engine='mineru', tier='advanced')
        waiting = self.poll('long.pdf', pending, lambda value: value.get('job_status') == 'waiting')
        self.assertEqual(waiting['status'], 'running')  # existing pollers keep waiting
        self.assertIn(waiting['parse']['state'], {'queued', 'parsing'})
        self.assertEqual(waiting['parse']['pages'], [1, 4])
        self.assertIn('Local MinerU', waiting['message'])
        self.assertIn('do not need to wait', waiting['message'])
        started = time.monotonic()
        quick = self.read('notes.md', wait=20, engine='builtin')
        elapsed = time.monotonic() - started
        self.assertTrue(quick['ok'], quick)
        self.assertEqual(quick['engine'], 'builtin')
        # The markdown read finished while the MinerU batch was still running.
        still = self.jobs.status(self.workspace, {'path': 'long.pdf', 'job_id': pending['job_id']})
        self.assertEqual(still['job_status'], 'waiting', f'markdown read took {elapsed:.2f}s')
        done = self.poll('long.pdf', pending, lambda value: value.get('job_status') in module.TERMINAL)
        self.assertTrue(done['ok'], done)
        self.assertEqual(done['engine'], 'mineru')
        self.assertEqual([unit['locator']['page'] for unit in done['units']], [1, 2, 3, 4])
        self.assertTrue(done['units'][0]['text'].startswith('PARSED long.pdf PAGE 1'))

    def test_one_persistent_worker_parses_every_batch_and_reads_hit_the_cache(self):
        value = self.jobs.prepare(self.workspace, {'paths': ['long.pdf'], 'engine': 'mineru', 'tier': 'basic'})
        self.assertTrue(value['ok'], value)
        self.assertEqual(value['documents'][0]['plan']['batches'], 3)  # 17 pages in 8-page basic batches
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline and not all(row['status'] == 'ready' for row in self.parses()):
            time.sleep(.05)
        self.assertEqual([row['status'] for row in self.parses()], ['ready'] * 3)
        self.assertEqual(len(self.starts()), 1, 'every batch must reuse the loaded worker')
        self.assertEqual([request['pages'] for request in self.requests()], [8, 8, 1])
        token, pages = None, []
        for _ in range(4):
            result = self.read('long.pdf', wait=20, engine='mineru', tier='basic', max_units=40, max_chars=100000,
                               **({'cursor': token} if token else {}))
            self.assertTrue(result['ok'], result)
            pages.extend(unit['locator']['page'] for unit in result['units'])
            token = result['next_cursor']
            if not token:
                break
        self.assertEqual(pages, list(range(1, 18)))
        self.assertEqual(len(self.requests()), 3, 'prepared batches are never parsed twice')
        progress = self.jobs.queue(self.workspace, {'paths': ['long.pdf']})
        self.assertEqual(progress['documents'][0]['state'], 'ready')
        self.assertEqual(progress['documents'][0]['tiers']['basic']['pages_ready'], 17)

    def test_read_ahead_parses_the_next_batches_without_being_asked(self):
        first = self.read('long.pdf', wait=20, engine='mineru', tier='advanced', max_units=40, max_chars=100000)
        self.assertTrue(first['ok'], first)
        self.assertEqual([unit['locator']['page'] for unit in first['units']], [1, 2, 3, 4])
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline and len(self.requests()) < 2:
            time.sleep(.05)
        rows = {(row['batch_start'], row['batch_end']): row for row in self.parses()}
        self.assertEqual(rows[(4, 8)]['origin'], 'read_ahead')
        self.assertEqual(rows[(4, 8)]['priority'], module.PARSE_PRIORITY['read_ahead'])
        # A continuing read (cursor) is sequential: two batches ahead are queued.
        second = self.read('long.pdf', wait=20, cursor=first['next_cursor'], max_units=40, max_chars=100000)
        self.assertTrue(second['ok'], second)
        self.assertEqual([unit['locator']['page'] for unit in second['units']], [5, 6, 7, 8])
        rows = {(row['batch_start'], row['batch_end']) for row in self.parses()}
        self.assertTrue({(8, 12), (12, 16)} <= rows, rows)

    def test_waiting_read_jumps_ahead_of_background_preparation(self):
        self.settings({'page_seconds': .1})
        prepared = self.jobs.prepare(self.workspace, {'paths': ['long.pdf'], 'engine': 'mineru', 'tier': 'advanced'})
        self.assertTrue(prepared['ok'], prepared)
        self.assertEqual(prepared['documents'][0]['plan']['batches'], 5)
        interactive = self.read('other.pdf', engine='mineru', tier='advanced')
        done = self.poll('other.pdf', interactive, lambda value: value.get('job_status') in module.TERMINAL)
        self.assertTrue(done['ok'], done)
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline and any(row['status'] != 'ready' for row in self.parses() if row['path'] == 'long.pdf'):
            time.sleep(.05)
        rows = self.parses()
        served = next(row['finished'] for row in rows if row['path'] == 'other.pdf' and row['batch_start'] == 0)
        before = sum(1 for row in rows if row['path'] == 'long.pdf' and row['finished'] < served)
        after = sum(1 for row in rows if row['path'] == 'long.pdf' and row['finished'] > served)
        self.assertEqual(before + after, 5)
        # The waiting read waits only for the batch already running, never for
        # the rest of the prepared document; background work then continues.
        self.assertLessEqual(before, 2)
        self.assertGreaterEqual(after, 3)

    def test_speed_is_measured_on_this_computer_and_used_for_estimates(self):
        self.settings({'page_seconds': .05})
        with self.jobs.connect() as db:
            self.assertEqual(self.jobs._speeds(db), {})
        first = self.read('long.pdf', wait=20, engine='mineru', tier='advanced')
        self.assertTrue(first['ok'], first)
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline and len(self.requests()) < 2:
            time.sleep(.05)
        with self.jobs.connect() as db:
            speeds = self.jobs._speeds(db)
            self.assertIn(('page', 'advanced'), speeds)
            self.jobs._enqueue_parse(db, None, 'prepare', key='f' * 64, scope='synthetic', source_hash='0' * 64,
                                     root=str(self.workspace), path='long.pdf', tier='advanced', batch_start=16,
                                     batch_end=17, unit_count=17)
            view = self.jobs._parse_progress(db)['f' * 64]
        self.assertEqual(view['state'], 'queued')
        self.assertIsNotNone(view['eta_seconds'])
        self.assertIn('on this computer', view['speed']['basis'])

    def test_stopping_a_waiting_read_stops_only_work_it_asked_for(self):
        self.settings({'page_seconds': .5})
        prepared = self.jobs.prepare(self.workspace, {'paths': ['other.pdf'], 'engine': 'mineru'})
        self.assertTrue(prepared['ok'], prepared)
        reading = self.read('long.pdf', engine='mineru')
        self.poll('long.pdf', reading, lambda value: value.get('job_status') == 'waiting')
        stopped = self.jobs.status(self.workspace, {'path': 'long.pdf', 'job_id': reading['job_id']}, cancel=True)
        self.assertEqual(stopped['job_status'], 'cancelled')
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline and any(row['path'] == 'long.pdf' and row['status'] in {'queued', 'running'} for row in self.parses()):
            time.sleep(.05)
        rows = self.parses()
        self.assertTrue(all(row['status'] in {'cancelled', 'ready'} for row in rows if row['path'] == 'long.pdf'))
        self.assertTrue(any(row['path'] == 'other.pdf' and row['status'] in {'queued', 'running', 'ready'} for row in rows),
                        'prepared background work keeps going')

    def test_auto_preparation_skips_documents_the_builtin_reader_covers(self):
        value = self.jobs.prepare(self.workspace, {'paths': ['long.pdf', 'notes.md', 'missing.pdf']})
        entries = {entry['path']: entry for entry in value['documents']}
        self.assertEqual(entries['long.pdf']['state'], 'not_needed')  # plain text pages only
        self.assertEqual(entries['notes.md']['state'], 'not_needed')
        self.assertFalse(entries['missing.pdf']['ok'])
        self.assertEqual(self.parses(), [])

    def test_auto_preparation_queues_documents_with_scans(self):
        import pymupdf
        from PIL import Image
        picture = self.root / 'scan.png'
        Image.new('RGB', (200, 120), 'white').save(picture)
        with pymupdf.open() as pdf:
            for number in range(1, 8):
                page = pdf.new_page()
                page.insert_text((50, 50), f'SCANNED PAGE {number}')
                if number == 5:
                    page.insert_image(pymupdf.Rect(50, 80, 250, 200), filename=str(picture))
            pdf.save(self.workspace / 'scanned.pdf')
        value = self.jobs.prepare(self.workspace, {'paths': ['scanned.pdf']})
        self.assertTrue(value['ok'], value)
        self.assertEqual(value['tier'], module.DEFAULT_AUTO_TIER)
        entry = value['documents'][0]
        self.assertTrue(entry['plan']['mineru_needed'])
        self.assertEqual(entry['plan']['batches'], 2)  # whole document, 6-page adaptive batches
        self.assertEqual({row['tier'] for row in self.parses()}, {'adaptive'})

    def test_queue_cancel_stops_background_parsing_and_lists_only_this_root(self):
        self.settings({'page_seconds': .5})
        self.jobs.prepare(self.workspace, {'paths': ['long.pdf'], 'engine': 'mineru'})
        listing = self.jobs.queue(self.workspace, {})
        self.assertEqual([entry['path'] for entry in listing['documents']], ['long.pdf'])
        elsewhere = self.root / 'elsewhere'
        elsewhere.mkdir()
        self.assertEqual(self.jobs.queue(elsewhere, {})['documents'], [])
        stopped = self.jobs.queue(self.workspace, {'paths': ['long.pdf'], 'cancel': True})
        self.assertIn(stopped['documents'][0]['state'], {'stopped', 'parsing'})
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline and any(row['status'] in {'queued', 'running'} for row in self.parses()):
            time.sleep(.05)
        self.assertTrue(all(row['status'] in {'cancelled', 'ready'} for row in self.parses()))
        self.assertIsNone(self.jobs.queue(self.workspace, {})['queue']['running'])

    def test_worker_crash_fails_the_waiting_read_and_a_fresh_read_retries(self):
        self.settings({'crash': True})
        crashed = self.read('other.pdf', engine='mineru')
        failed = self.poll('other.pdf', crashed, lambda value: value.get('job_status') in module.TERMINAL)
        self.assertFalse(failed['ok'])
        self.assertTrue(failed['error']['retryable'])
        self.assertIn('No remote service', failed['error']['message'])
        self.join_lanes()
        self.settings({})
        retry = self.read('other.pdf', wait=20, engine='mineru')
        self.assertTrue(retry['ok'], retry)
        self.assertGreaterEqual(len(self.starts()), 2)

    def test_local_decoding_slot_setting_reaches_the_worker(self):
        (self.root / 'runtime/document-jobs/mineru-local.json').write_text(json.dumps({'vlm_parallel': 3, 'vlm_device': 'cpu'}), encoding='utf-8')
        result = self.read('other.pdf', wait=20, engine='mineru')
        self.assertTrue(result['ok'], result)
        arguments = json.loads(self.starts()[0])['arguments']
        self.assertEqual(arguments[arguments.index('--vlm-parallel') + 1], '3')
        self.assertEqual(arguments[arguments.index('--vlm-device') + 1], 'cpu')
        (self.root / 'runtime/document-jobs/mineru-local.json').write_text(json.dumps({'vlm_device': 'vulkan'}), encoding='utf-8')
        self.assertEqual(self.jobs.local_settings(), {})  # diagnostics-only value is not a user setting
        (self.root / 'runtime/document-jobs/mineru-local.json').write_text(json.dumps({'vlm_parallel': 9}), encoding='utf-8')
        self.assertEqual(self.jobs.local_settings(), {})

    def test_word_images_are_parsed_and_attached_without_changing_text_or_cursor(self):
        from docx import Document
        from PIL import Image
        picture = self.root / 'figure.png'
        Image.new('RGB', (120, 80), 'white').save(picture)
        document = Document()
        document.add_paragraph('Synthetic paragraph before the picture.')
        document.add_picture(str(picture))
        document.add_paragraph('Synthetic paragraph after the picture.')
        document.save(self.workspace / 'report.docx')
        builtin = self.read('report.docx', wait=20, engine='builtin', ocr='off')
        result = self.read('report.docx', wait=20, engine='mineru')
        self.assertTrue(result['ok'], result)
        self.assertEqual(result['structure_scope'], 'embedded_images')
        self.assertEqual(result['image_analysis']['images'], 1)
        self.assertEqual([unit['text'] for unit in result['units']], [unit['text'] for unit in builtin['units']])
        self.assertEqual(result['next_cursor'], builtin['next_cursor'])
        pictured = next(unit for unit in result['units'] if unit.get('image_analysis'))
        self.assertEqual(pictured['image_analysis'][0]['text'], 'PARSED IMAGE CONTENT')
        self.assertEqual(pictured['status'], 'ready')
        self.assertNotIn('embedded_image_text_unverified', pictured['coverage']['unresolved'])
        rows = self.parses()
        self.assertEqual(len(rows), 1)
        self.assertNotEqual(rows[0]['source_hash'], rows[0]['owner_hash'])  # keyed by image content
        progress = self.jobs.queue(self.workspace, {'paths': ['report.docx']})
        self.assertEqual(progress['documents'][0]['state'], 'ready')  # grouped under the Word file
        again = self.read('report.docx', wait=20, engine='mineru', max_chars=600)
        self.assertTrue(again['ok'], again)
        self.assertEqual(len(self.requests()), 1, 'the same image is parsed once')

    def test_image_file_text_from_mineru_replaces_the_empty_builtin_unit(self):
        from PIL import Image
        Image.new('RGB', (160, 90), 'white').save(self.workspace / 'scan.png')
        result = self.read('scan.png', wait=20, engine='mineru')
        self.assertTrue(result['ok'], result)
        self.assertEqual(result['engine'], 'mineru')
        self.assertEqual(result['units'][0]['text'], 'PARSED IMAGE CONTENT')
        self.assertEqual(result['units'][0]['locator'], {'image': 1})

    def test_cache_key_format_is_unchanged(self):
        expected = hashlib.sha256(module._dump(['scope', '0' * 64, 'advanced', 0, 4, 'physical-pdf-batch-v2',
                                                hashlib.sha256(self.helper.read_bytes()).hexdigest()]).encode()).hexdigest()
        self.assertEqual(module.parse_key('scope', '0' * 64, 'advanced', 0, 4, self.helper), expected)


if __name__ == '__main__':
    unittest.main()
