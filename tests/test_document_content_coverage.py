"""Deterministic source fidelity fixtures; no private documents or model APIs."""
from __future__ import annotations

import base64
import importlib.util
import io
from pathlib import Path
import socket
import tempfile
import unittest
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("coverage_reader", ROOT / "scripts/workspace_document_reader.py")
reader = importlib.util.module_from_spec(spec)
spec.loader.exec_module(reader)


class ContentCoverageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="document-fidelity-fixture-", dir=ROOT / "runtime")
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def read(self, path, **options):
        result = reader.handle_request({"root": str(self.root), "path": path, "ocr": "off", **options})
        self.assertTrue(result["ok"], result)
        return result

    def image(self, detailed=False):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new("RGB", (1400, 420), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.load_default(size=64 if detailed else 42)
        draw.text((60, 60), "Local document reading" if detailed else "Synthetic scanned region", fill="black", font=font)
        draw.text((60, 160), "Synthetic scan text" if detailed else "Figure A = 123", fill="black", font=font)
        stream = io.BytesIO()
        image.save(stream, format="PNG")
        return stream.getvalue()

    def math(self, body):
        from lxml import etree
        return etree.fromstring(f'<m:oMath xmlns:m="{reader.MATH_NS}">{body}</m:oMath>')

    def test_word_equation_table_and_image_keep_source_identity(self):
        from docx import Document
        document = Document()
        paragraph = document.add_paragraph("Energy: ")
        paragraph._p.append(self.math('<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f><m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>'))
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Variable"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "root"
        table.cell(1, 1).paragraphs[0]._p.append(self.math('<m:rad><m:deg><m:r><m:t>3</m:t></m:r></m:deg><m:e><m:r><m:t>8</m:t></m:r></m:e></m:rad>'))
        document.add_picture(io.BytesIO(self.image()))
        document.sections[0].header.paragraphs[0].text = "Synthetic header"
        document.sections[0].footer.paragraphs[0].text = "Synthetic footer"
        path = self.root / "math.docx"
        document.save(path)
        before = path.read_bytes()
        result = self.read("math.docx")
        self.assertIn(r"\frac{1}{2}{x}^{2}", result["text"])
        self.assertIn(r"\sqrt[3]{8}", result["units"][1]["table"]["rows"][1][1])
        self.assertEqual(result["coverage"]["formula_count"], 2)
        self.assertEqual(result["coverage"]["table_count"], 1)
        self.assertIn("Synthetic header", result["text"])
        self.assertIn("Synthetic footer", result["text"])
        self.assertEqual(result["units"][0]["equations"][0]["status"], "complete")
        self.assertEqual(result["units"][2]["status"], "needs_ocr")
        image_ref = result["units"][2]["assets"][0]
        image = self.read("math.docx", action="asset", asset_id=image_ref["id"], document_sha256=result["document_sha256"])
        self.assertTrue(base64.b64decode(image["asset"]["base64"]).startswith(b"\x89PNG"))
        self.assertEqual(path.read_bytes(), before)

    def test_unknown_equation_structure_is_not_falsely_complete(self):
        from docx import Document
        document = Document()
        document.add_paragraph()._p.append(self.math('<m:phant><m:e><m:r><m:t>x</m:t></m:r></m:e></m:phant>'))
        document.save(self.root / "unsupported.docx")
        result = self.read("unsupported.docx")
        self.assertEqual(result["status"], "partial")
        self.assertIn("x", result["text"])
        self.assertEqual(result["units"][0]["equations"][0]["status"], "partial")
        self.assertIn("unhandled_omml_phant", result["coverage"]["unresolved"])

    def test_word_mixed_text_does_not_mask_unrecognized_embedded_images(self):
        from docx import Document
        document = Document()
        paragraph = document.add_paragraph("Only the caption is native text")
        paragraph.add_run().add_picture(io.BytesIO(self.image()))
        document.save(self.root / "caption.docx")
        with mock.patch.object(reader.Document, "local_ocr", return_value=""):
            result = self.read("caption.docx", ocr="auto")
        self.assertEqual(result["status"], "partial")
        self.assertIn("embedded_image_text_unverified", result["coverage"]["unresolved"])
        self.assertEqual(result["coverage"]["ocr_units"], 0)
        self.assertTrue(result["units"][0]["assets"])

    def test_matrix_limits_literal_tex_not_executed(self):
        matrix = self.math('<m:m><m:mr><m:e><m:r><m:t>1</m:t></m:r></m:e><m:e><m:r><m:t>2</m:t></m:r></m:e></m:mr><m:mr><m:e><m:r><m:t>3</m:t></m:r></m:e><m:e><m:r><m:t>4</m:t></m:r></m:e></m:mr></m:m>')
        result = reader.omml_equation(matrix)
        self.assertEqual(result["text"], r"\begin{matrix}1 & 2 \\ 3 & 4\end{matrix}")
        equation = self.math('<m:r><m:t>\\input{secrets}</m:t></m:r>')
        result = reader.omml_equation(equation)
        self.assertNotIn(r"\input{", result["text"])
        self.assertIn(r"\backslash", result["text"])

    def test_word_content_control_is_read_not_silently_skipped(self):
        from docx import Document
        from docx.oxml import OxmlElement
        document = Document()
        paragraph = document.add_paragraph("Inside a Word content control")
        document.element.body.remove(paragraph._p)
        control, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        content.append(paragraph._p)
        control.append(content)
        document.element.body.insert(0, control)
        document.save(self.root / "control.docx")
        result = self.read("control.docx")
        self.assertIn("Inside a Word content control", result["text"])
        self.assertEqual(result["units"][0]["locator"], {"paragraph": 1})

    def test_partial_read_does_not_claim_whole_document_or_bypass_budget(self):
        from docx import Document
        document = Document()
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "a" * 800
        table.cell(1, 0).text = "b" * 800
        document.save(self.root / "long.docx")
        first = self.read("long.docx", max_chars=256)
        self.assertFalse(first["coverage"]["document_fully_read"])
        self.assertNotIn("table", first["units"][0])
        self.assertEqual(first["units"][0]["coverage"]["table_structure"], "fragmented")
        cursor, chunks = first["next_cursor"], [first["units"][0]["text"]]
        while cursor:
            result = self.read("long.docx", max_chars=256, cursor=cursor)
            self.assertFalse(result["coverage"]["document_fully_read"])
            chunks.extend(row["text"] for row in result["units"])
            cursor = result["next_cursor"]
        self.assertEqual("".join(chunks), "a" * 800 + "\n" + "b" * 800)

    def test_mixed_pdf_header_does_not_hide_scanned_body(self):
        import pymupdf
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((30, 30), "Header only")
        page.insert_image(pymupdf.Rect(20, 100, 580, 270), stream=self.image())
        page.draw_circle((300, 420), 60)
        document.save(self.root / "mixed.pdf")
        document.close()
        with mock.patch.object(reader.Document, "local_ocr", return_value="Recognized scan body") as ocr:
            result = self.read("mixed.pdf", ocr="auto")
        self.assertEqual(ocr.call_count, 1)
        self.assertIn("Header only", result["text"])
        self.assertIn("Recognized scan body", result["text"])
        self.assertEqual(result["units"][0]["extraction"], "text+ocr")
        self.assertEqual(result["coverage"]["ocr_units"], 1)
        self.assertTrue(result["coverage"]["visual_review_required"])
        self.assertEqual(result["coverage"]["semantic_completeness"], "not_guaranteed")
        no_ocr = self.read("mixed.pdf")
        self.assertEqual(no_ocr["status"], "partial")
        self.assertIn("raster_or_vector_text_unverified", no_ocr["coverage"]["unresolved"])

    def test_visual_only_pdf_retains_page_without_fake_formula_text(self):
        import pymupdf
        document = pymupdf.open()
        page = document.new_page()
        page.draw_line((40, 40), (160, 40))
        page.draw_circle((100, 100), 35)
        document.save(self.root / "visual.pdf")
        document.close()
        result = self.read("visual.pdf")
        self.assertEqual(result["status"], "needs_ocr")
        self.assertEqual(result["text"], "")
        self.assertEqual(result["coverage"]["visual_units"], 1)
        self.assertEqual(result["coverage"]["formula_count"], 0)
        self.assertTrue(result["units"][0]["assets"])

    def test_real_offline_english_scan_with_packaged_ocr(self):
        try:
            import rapidocr
        except ImportError:
            self.skipTest("Run this test in the portable packaged OCR runtime")
        import pymupdf
        document = pymupdf.open()
        page = document.new_page(width=700, height=300)
        page.insert_image(pymupdf.Rect(0, 0, 700, 210), stream=self.image(detailed=True))
        document.save(self.root / "english-scan.pdf")
        document.close()
        with mock.patch.object(socket.socket, "connect", side_effect=AssertionError("no network allowed")), mock.patch.object(socket, "create_connection", side_effect=AssertionError("no network allowed")):
            result = self.read("english-scan.pdf", ocr="auto")
        self.assertEqual(result["units"][0]["extraction"], "ocr", result)
        self.assertIn("Local document reading", result["text"])
        self.assertIn("Synthetic scan text", result["text"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
