"""Small original synthetic artifacts for actual parser regressions."""
from pathlib import Path


def create_fixtures(root):
    from reportlab.pdfgen import canvas
    from docx import Document
    from pptx import Presentation
    from openpyxl import Workbook
    root=Path(root);root.mkdir(parents=True,exist_ok=True)
    pdf=canvas.Canvas(str(root/'fixture.pdf'));pdf.drawString(60,750,'Synthetic PDF page 1');pdf.showPage();pdf.save()
    doc=Document();doc.add_heading('Synthetic Word',level=1);doc.save(root/'fixture.docx')
    deck=Presentation()
    for _ in range(2):
        slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='Synthetic PPT'
    deck.save(root/'fixture.pptx')
    book=Workbook();book.active['A1']='Synthetic sheet';book.save(root/'fixture.xlsx');book.close()
    (root/'fixture.md').write_text('# Synthetic Markdown\n',encoding='utf-8')
