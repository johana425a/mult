"""Isolated OpenMAIC native-reader bridge; synthetic documents only."""
from __future__ import annotations

import base64
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
PYTHON = Path(os.environ.get('DEEPSTEAM_TEST_PYTHON', sys.executable))
OCR_PYTHON = PYTHON
WRAPPER = ROOT / 'scripts/openmaic_document_extract.py'


class OpenMAICDocumentReaderTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='openmaic-reader-test-')
        self.root = Path(self.temporary.name).resolve()
        self.work = self.root / 'extract-fixture'
        self.work.mkdir()

    def tearDown(self):
        self.temporary.cleanup()

    def run_reader(self, source, python=PYTHON, **extra):
        allowed = {'SYSTEMROOT', 'WINDIR', 'COMSPEC', 'PATH', 'PATHEXT', 'TEMP', 'TMP'}
        environment = {key: value for key, value in os.environ.items() if key.upper() in allowed}
        environment['DEEPSTEAM_DOCUMENT_TEMP_ROOT'] = str(self.root)
        result = subprocess.run([str(python), '-I', '-B', str(WRAPPER)],
                                input=json.dumps({'source': str(source), **extra}).encode(),
                                stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                env=environment, timeout=90, check=False,
                                **({'creationflags': 0x08000000} if os.name == 'nt' else {}))
        return result.returncode, json.loads(result.stdout)

    def test_word_preserves_header_locations_and_images(self):
        from docx import Document
        from PIL import Image
        document = Document()
        document.add_paragraph('Synthetic original body')
        header = document.sections[0].header.paragraphs[0]
        header.text = 'Synthetic header text'
        image = io.BytesIO()
        Image.new('RGB', (20, 20), 'blue').save(image, format='PNG')
        header.add_run().add_picture(io.BytesIO(image.getvalue()))
        source = self.work / 'source.docx'
        document.save(source)
        code, result = self.run_reader(source)
        self.assertEqual(code, 0, result)
        self.assertTrue(result['coverage']['document_fully_read'])
        supplement = next(unit for unit in result['units'] if unit['locator'].get('part'))
        self.assertEqual(supplement['locator'], {'part': '/word/header1.xml', 'section': 'word_supplement'})
        self.assertIn('Synthetic header text', supplement['text'])
        image = next(image for image in result['images'] if image['locator'].get('part'))
        self.assertTrue(base64.b64decode(image['base64']).startswith(b'\x89PNG'))
        self.assertNotIn('data', image)

    def test_pdf_text_uses_physical_pages_and_complete_units(self):
        import pymupdf
        document = pymupdf.open()
        for number in range(3):
            page = document.new_page()
            page.insert_text((40, 40), f'Synthetic page {number + 1}')
        source = self.work / 'source.pdf'
        document.save(source)
        document.close()
        code, result = self.run_reader(source)
        self.assertEqual(code, 0, result)
        self.assertEqual(len(result['units']), 3)
        self.assertEqual(result['units'][2]['locator'], {'page': 3, 'page_kind': 'pdf_physical'})
        self.assertEqual(result['coverage']['semantic_completeness'], 'not_guaranteed')

    def test_presentation_retains_original_slide_locator(self):
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = 'Synthetic slide title'
        slide.placeholders[1].text = 'Synthetic original paragraph'
        source = self.work / 'source.pptx'
        document.save(source)
        code, result = self.run_reader(source)
        self.assertEqual(code, 0, result)
        self.assertEqual(result['units'][0]['locator'], {'slide': 1})
        self.assertIn('Synthetic original paragraph', result['units'][0]['text'])

    def test_rejects_source_outside_private_copy_and_caller_root(self):
        external = self.root / 'private.docx'
        external.write_bytes(b'synthetic-invalid-file')
        code, result = self.run_reader(external)
        self.assertNotEqual(code, 0)
        self.assertFalse(result['ok'])
        self.assertNotIn(str(external), result['error'])
        code, result = self.run_reader(external, root=str(self.root))
        self.assertNotEqual(code, 0)
        self.assertFalse(result['ok'])
        self.assertTrue(external.is_file())

    @unittest.skipUnless(OCR_PYTHON.is_file(), 'packaged OCR runtime unavailable')
    def test_packaged_python_recognizes_synthetic_scan_without_a_provider(self):
        import pymupdf
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new('RGB', (1100, 400), 'white')
        ImageDraw.Draw(image).text((40, 80), 'NATIVE UPLOAD TEST 24680', fill='black',
                                  font=ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', 50))
        content = io.BytesIO()
        image.save(content, format='PNG')
        document = pymupdf.open()
        page = document.new_page(width=700, height=300)
        page.insert_image(page.rect, stream=content.getvalue())
        source = self.work / 'source.pdf'
        document.save(source)
        document.close()
        code, result = self.run_reader(source, python=OCR_PYTHON)
        self.assertEqual(code, 0, result)
        self.assertIn('24680', '\n'.join(unit['text'] for unit in result['units']))
        self.assertTrue(result['images'])


if __name__ == '__main__':
    unittest.main()
