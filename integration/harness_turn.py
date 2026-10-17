"""Native teaching is a tool provider; Harness owns the conversation loop.

The registry is request scoped, authenticated and bound to the original user's
ContextVars. Historical messages are read-only context, never replayed actions.
"""
from __future__ import annotations
import asyncio
import contextvars
import hashlib
import json
import os
import time
import uuid
from pathlib import Path
from fastapi import Request, HTTPException

_contexts = {}
_record_locks = {}
_health = (0.0, False)


def _source_snapshot(context, owner):
    """Freeze authoritative source versions, not the current viewport or prompt."""
    from integration.execution_workspace_binding import owner_workspace_binding
    workspace_id = context.runtime.workspace.workspace_id
    grant = owner_workspace_binding(owner, workspace_id)
    metadata = context.metadata
    digest = lambda value: hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
        separators=(',', ':')).encode('utf-8')).hexdigest()
    rows = [{'owner':owner, 'scope':workspace_id, 'source_id':'workspace-grant', 'version':digest(grant)}]
    for name in ('source_index', 'mastery_topic_source_index'):
        for source_id, content in sorted((metadata.get(name) or {}).items()):
            rows.append({'owner':owner, 'scope':workspace_id, 'source_id':name+':'+source_id,
                'version':digest(content), 'sha256':digest(content)})
    material_ids = {metadata.get('reading_material_id')} - {None, ''}
    for name in ('reading_references', 'source_references'):
        for reference in metadata.get(name) or []:
            if isinstance(reference, dict) and reference.get('material_id'):
                material_ids.add(reference['material_id'])
    if material_ids:
        from deeptutor.reading.store import ReadingStore
        store = ReadingStore()
        for material_id in sorted(material_ids):
            manifest = store.manifest(material_id)
            requested = metadata.get('reading_material_revision') if material_id == metadata.get('reading_material_id') else None
            if requested and requested != manifest.revision:
                raise HTTPException(409, {'code':'execution_source_context_changed', 'message':'The selected material revision changed.'})
            rows.append({'owner':owner, 'scope':workspace_id, 'source_id':'reading:'+material_id,
                'version':manifest.revision, 'sha256':manifest.source_hash})
    return rows


async def _release_context_when_settled(ref, entry):
    from integration.harness_execution import ExecutionClient
    from urllib.parse import urlencode
    task = entry.get('harness_task')
    if not task:
        _contexts.pop(ref, None)
        return
    while _contexts.get(ref) is entry:
        await asyncio.sleep(20)
        try:
            value = await ExecutionClient().request('GET', '/' + task['id'] + '?' + urlencode({'owner': entry['owner']}))
            if value['task']['status'] in {'completed', 'failed', 'cancelled', 'interrupted'}:
                _contexts.pop(ref, None)
                return
        except Exception:
            continue


def enabled():
    return os.environ.get('DEEPSTEAM_HARNESS_CORE') == '1'


def execution_available(kind='shell'):
    global _health
    if not enabled() or kind not in {'shell', 'python', 'node'}:
        return False
    if time.monotonic() - _health[0] > 3:
        import urllib.request
        try:
            url = os.environ.get('DEEPSTEAM_HARNESS_ORIGIN', 'http://127.0.0.1:3796') + '/health'
            with urllib.request.build_opener(urllib.request.ProxyHandler({})).open(url, timeout=1) as response:
                okay = bool(json.load(response).get('healthy'))
        except Exception:
            okay = False
        _health = (time.monotonic(), okay)
    return _health[1]


def validate_file(path):
    """Validate bytes before native presentation; no empty-directory success."""
    path = Path(path)
    if os.name == 'nt' and path.is_absolute() and not str(path).startswith('\\\\?\\'):
        path = Path('\\\\?\\' + str(path))
    if not path.is_file() or not path.stat().st_size:
        raise ValueError('Output file is missing or empty: ' + path.name)
    ext = path.suffix.lower()
    if ext == '.json':
        json.loads(path.read_text(encoding='utf-8-sig'))
    elif ext in {'.md', '.txt', '.csv'}:
        path.read_text(encoding='utf-8-sig')
    elif ext == '.docx':
        from docx import Document
        Document(path)
    elif ext == '.pdf':
        from pypdf import PdfReader
        if not PdfReader(path).pages:
            raise ValueError('PDF has no pages')
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True); book.close()
    elif ext == '.pptx':
        from pptx import Presentation
        Presentation(path)
    with path.open('rb') as file:
        digest = hashlib.file_digest(file, 'sha256').hexdigest()
    return {'filename': path.name, 'format': ext.lstrip('.'), 'size_bytes': path.stat().st_size, 'sha256': digest, 'verified': True}


def install_routes(app, authenticate):
    from fastapi import Request, HTTPException

    @app.post('/api/contest/internal/harness-surface-binding')
    async def surface_binding(request: Request):
        authenticate(request)
        body = await request.json()
        from integration.execution_workspace_binding import native_execution_owner
        from deeptutor.multi_user.models import LOCAL_ADMIN_ID
        from deeptutor.services.auth import AUTH_ENABLED
        from deeptutor.multi_user.paths import user_context
        from deeptutor.services.workspace import get_content_workspace_service
        from deeptutor.services.model_selection.runtime import resolve_model_selection
        surface = body.get('surface')
        if surface not in {'openmaic-workbench', 'openmaic-classroom'}:
            raise HTTPException(422, {'code': 'surface_invalid', 'message': 'Unknown interaction surface.'})
        if AUTH_ENABLED:
            raise HTTPException(403, {'code':'classroom_owner_binding_required',
                'message':'This classroom bridge requires an authenticated account binding in multi-user mode.'})
        user = native_execution_owner(LOCAL_ADMIN_ID)
        with user_context(user):
            try:
                _, concrete, snapshot = resolve_model_selection()
            except Exception as exc:
                from deeptutor.services.llm.request_compat import logged_error_text
                raise HTTPException(409, {'code':'model_selection_unavailable', 'phase':'model_binding',
                    'message':logged_error_text(exc)[:1200], 'trace_id':uuid.uuid4().hex, 'retryable':False}) from None
            try:
                runtime = get_content_workspace_service().create_runtime_context(capability=surface,
                    session_id=str(body.get('conversation_id') or 'classroom'),
                    turn_id=str(body.get('turn_id') or uuid.uuid4().hex))
            except OSError as exc:
                raise HTTPException(500, {'code':'workspace_output_unavailable','phase':'workspace_binding',
                    'message':str(exc)[:1200],'trace_id':uuid.uuid4().hex,'retryable':False}) from None
            return {'owner': user.id, 'native_workspace_id': runtime.workspace_id,
                'output_dir': str(runtime.output_dir), 'llm_selection': concrete,
                'resolved_model': snapshot}

    @app.post('/api/contest/internal/surface-tools/invoke')
    async def invoke(request: Request):
        authenticate(request)
        body = await request.json()
        if body.get('surface') == 'learning-records':
            return await _learning_record(body)
        entry = _contexts.get(body.get('context_ref'))
        if not entry or body.get('owner') != entry['owner'] or body.get('surface') != entry['surface']:
            raise HTTPException(409, {'code': 'surface_context_expired', 'message': 'The teaching context is no longer active. Continue in the original conversation to reconnect it.'})
        name = body.get('name')
        args = body.get('arguments')
        if name not in entry['names'] or not isinstance(args, dict) or any(k.startswith('_') for k in args):
            raise HTTPException(422, {'code': 'surface_tool_invalid', 'message': 'Use a registered tool and its public arguments.'})
        call_id = body.get('call_id') or uuid.uuid4().hex
        # Repeated HTTP delivery observes the same operation, never writes twice.
        if call_id not in entry['calls']:
            entry['calls'][call_id] = entry['vars'].run(asyncio.create_task, _invoke(entry, name, args, call_id))
        try:
            return await asyncio.shield(entry['calls'][call_id])
        except Exception as exc:
            from deeptutor.services.llm.request_compat import logged_error_text
            raise HTTPException(500, {'code': 'surface_tool_failed', 'phase': 'teaching_tool',
                'message': logged_error_text(exc)[:1600], 'trace_id': call_id, 'retryable': False}) from None


async def _learning_record(body):
    args = body.get('arguments') or {}
    key = (body.get('owner'), args.get('workspace_id'))
    lock = _record_locks.setdefault(key, asyncio.Lock())
    async with lock:
        return await _save_learning_record(body)


def _record_operation(body):
    args = body.get('arguments') or {}
    operation = body.get('operation')
    if operation is not None:
        if not isinstance(operation, dict) or not all(isinstance(operation.get(key), str) and operation[key]
                for key in ('task_id', 'agent_id', 'call_id')) or operation['call_id'] != body.get('call_id'):
            raise HTTPException(422, {'code':'record_operation_invalid'})
        identity = ['harness-v1', body['owner'], args['workspace_id'], operation['task_id'], operation['agent_id'], operation['call_id']]
    else:
        # Compatibility requests are conversation-scoped. Old bare harness_calls
        # remain audit metadata and are never treated as global write receipts.
        if not args.get('conversation_id') or not body.get('call_id'):
            raise HTTPException(422, {'code':'record_operation_required'})
        identity = ['legacy-v2', body['owner'], args['workspace_id'], args['conversation_id'], args.get('source'), body['call_id']]
    encode = lambda value: json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
    return {'operation_id':hashlib.sha256(encode(identity)).hexdigest(),
        'fingerprint':hashlib.sha256(encode(args)).hexdigest(), **(operation or {})}


async def _save_learning_record(body):
    from integration.execution_workspace_binding import owner_workspace_binding, native_execution_owner
    from deeptutor.multi_user.paths import user_context
    from deeptutor.services.notebook import get_notebook_manager
    from deeptutor.services.workspace.context import workspace_context
    args = body.get('arguments') or {}
    owner_workspace_binding(body['owner'], args.get('workspace_id'))
    user = native_execution_owner(body['owner'])
    # Service-to-service requests do not inherit the original renderer header.
    # Resolve both stores under the admitted workspace, after binding its owner.
    with user_context(user), workspace_context(args['workspace_id']):
        manager = get_notebook_manager()
        notebook = manager.get_or_create_notebook('Learning records · ' + args['workspace_id'],
            description='Conversation archives and learning summaries. These records do not assert mastery.')
        records = manager.get_records(notebook['id'])
        if args.get('action') == 'list':
            return {'notebook_id': notebook['id'], 'records': records}
        operation = _record_operation(body)
        for record in records:
            for previous in (record.get('metadata') or {}).get('harness_operations', []):
                if previous.get('operation_id') == operation['operation_id']:
                    if previous.get('fingerprint') != operation['fingerprint']:
                        raise HTTPException(409, {'code':'record_operation_conflict', 'message':'This operation already saved different content.'})
                    return {'notebook_id': notebook['id'], 'record': record, 'reused': True}
        content = str(args.get('content') or '')
        archive_snapshot = None
        if args.get('action') == 'archive_chat':
            if str(args.get('conversation_id') or '').startswith('unified_'):
                from deeptutor.services.session import get_session_store
                history = await get_session_store().get_messages(args['conversation_id'])
            else:
                history = json.loads(args.get('history') or '[]')
            if not isinstance(history, list) or not history:
                raise HTTPException(409, {'code':'chat_history_unavailable',
                    'message':'No original messages were found in this conversation workspace. No archive was saved.'})
            # Preserve the returned message objects alongside the readable text
            # projection, including their parts/attachments/tool fields. This is
            # a snapshot of the message service response, not all UI/runtime events.
            from deeptutor.services.workspace import get_content_workspace_service
            from scripts.repair_bundle import atomic
            from scripts.transactional_upgrade import longpath
            archive_data = json.dumps({'schema':1,'conversation_id':args.get('conversation_id'),
                'scope':'message_service_response' if str(args.get('conversation_id') or '').startswith('unified_') else 'supplied_classroom_messages',
                'messages':history},ensure_ascii=False,separators=(',',':')).encode('utf-8')
            archive_hash = hashlib.sha256(archive_data).hexdigest()
            relative = Path('outputs/chat-archives')/(archive_hash+'.json')
            binding = get_content_workspace_service().binding_by_id(args['workspace_id'])
            target = longpath(binding.root/relative)
            if not target.exists():
                atomic(target,archive_data)
            if hashlib.sha256(target.read_bytes()).hexdigest()!=archive_hash:
                raise ValueError('CHAT_ARCHIVE_SNAPSHOT_CHANGED')
            archive_snapshot = {'path':relative.as_posix(),'sha256':archive_hash,'bytes':len(archive_data),
                'display_kind':'text_projection','scope':'message_objects_only','message_count':len(history)}
            chunks = []
            for row in history:
                if row.get('role') not in {'user', 'assistant'}:
                    continue
                value = row.get('content', row.get('parts', ''))
                if isinstance(value, list):
                    value = '\n'.join(p.get('text', '') for p in value if p.get('type') == 'text')
                chunks.append('## ' + row['role'] + '\n\n' + str(value))
            content = '\n\n'.join(chunks)
            if not content.strip():
                content = 'Conversation message snapshot: ' + str(len(history)) + ' messages. See the linked JSON snapshot for non-text content.'
        if not content.strip():
            raise HTTPException(422, {'code':'record_empty','message':'No content was supplied for this record.'})
        if args.get('action') == 'append':
            record = next((r for r in records if r['id'] == args.get('record_id')), None)
            if not record:
                raise HTTPException(404, {'code':'record_missing','message':'The original record was not found.'})
            if (record.get('metadata') or {}).get('record_kind') == 'chat_archive':
                raise HTTPException(422, {'code':'archive_is_source_snapshot','message':'Conversation archives contain actual messages only. Save an added summary as a separate learning_summary record.'})
            metadata = {**(record.get('metadata') or {})}
            metadata['harness_calls'] = [*metadata.get('harness_calls', []), body['call_id']]
            metadata['harness_operations'] = [*metadata.get('harness_operations', []), operation]
            manager.update_record(notebook['id'],record['id'],output=record.get('output','')+'\n\n'+content,metadata=metadata)
            return {'notebook_id':notebook['id'],'record_id':record['id'],'updated':True}
        if args.get('action') not in {'save_summary','archive_chat'}:
            raise HTTPException(422, {'code':'record_action_invalid','message':'Unknown record action.'})
        result = manager.add_record([notebook['id']], 'chat', str(args.get('title') or 'Learning record'), '', content,
            metadata={'record_kind':'chat_archive' if args['action']=='archive_chat' else 'learning_summary',
                'workspace_id':args['workspace_id'],'conversation_id':args.get('conversation_id'),
                'source':args.get('source'),'source_references':args.get('source_references',[]),
                'harness_calls':[body['call_id']], 'harness_operations':[operation], 'captured_at':time.time(), 'mastery_asserted':False,
                **({'archive_snapshot':archive_snapshot} if archive_snapshot else {})})
        return {'notebook_id':notebook['id'],**result,**({'archive_snapshot':archive_snapshot} if archive_snapshot else {})}


async def _invoke(entry, name, args, call_id):
    pipeline, context, stream = entry['pipeline'], entry['context'], entry['stream']
    if 'source_snapshot' in entry and _source_snapshot(context, entry['owner']) != entry['source_snapshot']:
        raise HTTPException(409, {'code':'execution_source_context_changed',
            'message':'The source or its authorization changed. Start a new task with the current source.'})
    if name in entry.get('deferred', set()):
        await pipeline._execute_tool_call('load_tools', pipeline._augment_tool_kwargs('load_tools', {'names':[name]}, context), stream=stream)
        entry['deferred'].discard(name)
    if name == 'workspace_present':
        from deeptutor.services.workspace import get_content_workspace_service
        service = get_content_workspace_service()
        binding = service.binding_by_id(context.runtime.workspace.workspace_id)
        for item in args.get('items', []):
            path = (Path(binding.root) / item['path']).resolve()
            if not path.is_relative_to(Path(binding.root).resolve()):
                raise ValueError('Presentation path must be in the current workspace')
            await asyncio.to_thread(validate_file, path)
    dispatch = await pipeline._dispatch_tool_calls(tool_calls=[{'id': call_id, 'name': name,
        'arguments': json.dumps(args, ensure_ascii=False)}], context=context, stream=stream,
        iteration_index=len(entry['calls']), stage='executing')
    if dispatch.pause:
        await pipeline._await_user_reply_and_resolve(context=context, stream=stream, dispatch=dispatch)
    # A successfully posted assessment card is this turn's final artefact.
    # Honour the same post-tool hook as AgentLoop before asking the model for
    # another round. The hook may consume its marker, so retain its decision.
    final_override = pipeline._capability_final_text_override(context, '')
    if final_override is not None:
        entry['terminal_override'] = final_override
    if name == 'workspace_present' and not any(
        value.get('workspace_items') for value in dispatch.tool_metadata_by_id.values()
        if isinstance(value, dict)
    ):
        detail = '\n'.join(str(m.get('content') or '') for m in dispatch.tool_messages)
        raise RuntimeError('FILE_DELIVERY_FAILED: ' + detail[:1500])
    sources = []
    for reference in dispatch.sources or []:
        reference = dict(reference)
        if reference.get('type') == 'reading' and reference.get('material_id'):
            from deeptutor.reading.store import ReadingStore
            store = ReadingStore()
            manifest = store.manifest(reference['material_id'])
            if manifest.revision == reference.get('material_revision'):
                aliases = store.registered_source_aliases(manifest.material_id)
                raw = store.raw_path(manifest.material_id)
                if raw:
                    aliases.append(str(raw.resolve()))
                material_key = f'reading:{manifest.material_id}@{manifest.revision}'
                reference['_source_binding'] = {'owner': entry['owner'], 'scope': context.runtime.workspace.workspace_id,
                    'source_id': aliases[0] if aliases else material_key, 'version': manifest.revision,
                    'sha256': manifest.source_hash, 'aliases': [*aliases, material_key]}
        sources.append(reference)
    return {'content': '\n'.join(str(m.get('content') or '') for m in dispatch.tool_messages),
        'sources': sources, 'model_messages': dispatch.model_messages,
        'metadata': dispatch.tool_metadata_by_id,
        'terminate': dispatch.terminate or final_override is not None}


async def run_pipeline(pipeline, context, stream):
    from deeptutor.agents.loop.agent_loop import LoopOutcome
    from deeptutor.multi_user.context import get_current_user
    from deeptutor.core.stream import StreamEvent, StreamEventType
    from integration.harness_execution import ExecutionClient, selection_for_config, stable_task_key
    from deeptutor.services.model_selection.runtime import resolved_model_snapshot
    if not execution_available():
        raise RuntimeError('HARNESS_UNAVAILABLE: The shared execution service is not running. Start the DeepSteam runtime and retry this turn.')
    await pipeline._prepare_deferred_tools(context)
    await pipeline._prepare_kb_manifests(context)
    pipeline._exec_enabled = True
    names = [n for n in pipeline._compose_enabled_tools(context) if n not in {'exec', 'execute_workspace_task', 'workspace_export'}]
    schemas = pipeline._build_llm_tool_schemas(names, context)
    if pipeline._tool_view is not None:
        pipeline._tool_view.attach(schemas)
    # Advertise scoped external tools without eagerly connecting every provider.
    deferred = set()
    for tool in pipeline._deferred_pool:
        definition = tool.get_definition()
        if definition.name in {'exec', 'execute_workspace_task', 'workspace_export'}:
            continue
        if definition.name not in names:
            names.append(definition.name)
            schemas.append(definition.to_openai_schema())
            deferred.add(definition.name)
    reading = context.active_capability in {'reading', 'watching'} or context.metadata.get('reading_material_id') or context.metadata.get('immersive_reading_mode')
    learning = context.metadata.get('mastery_path_id') or context.metadata.get('mastery_mode') or context.active_capability in {'mastery', 'mastery_path', 'course'}
    surface = 'reading' if reading else 'learning' if learning else 'chat'
    seed = pipeline._capability_pre_loop_seed(context)
    # These hooks fetch existing state only. The optional explore_context hook
    # starts another agent loop, so discovery now belongs to Harness itself.
    for capability in pipeline._active_loop_capabilities(context):
        hook = getattr(capability, 'pre_loop', None)
        if callable(hook) and capability.__class__.__module__.split('.')[-2] != 'explore_context':
            block = await hook(context, stream, usage=pipeline._usage)
            if block and getattr(block, 'content', ''):
                seed += '\n\n' + block.content
    messages = pipeline._build_loop_messages(context=context, enabled_tools=names, kb_seed=seed)
    images = []
    def textual(message, *, current=False):
        value = message.get('content', '')
        if not isinstance(value, list):
            return value
        text = []
        for part in value:
            if part.get('type') == 'text':
                text.append(part.get('text', ''))
            elif part.get('type') == 'image_url':
                if current:
                    images.append(part['image_url']['url'])
                text.append('[Source image attached to this turn]' if current else '[Historical image; retrieve original if needed]')
        return '\n'.join(text)
    system = '\n\n'.join(str(m.get('content', '')) for m in messages if m.get('role') == 'system')
    history = [{'role': m['role'], 'content': textual(m)} for m in messages[1:-1] if m.get('role') in {'user', 'assistant'}]
    current_messages = [{**m, 'content': textual(m, current=True)} for m in messages[-1:]]
    owner = get_current_user().id
    ref = uuid.uuid4().hex
    snapshot = _source_snapshot(context, owner)
    entry = {'pipeline': pipeline, 'context': context, 'stream': stream, 'owner': owner, 'source_snapshot': snapshot,
        'surface': surface, 'names': set(names), 'deferred': deferred, 'vars': contextvars.copy_context(), 'calls': {}}
    _contexts[ref] = entry
    interaction = {'surface': surface, 'context_ref': ref, 'source_snapshot': snapshot, 'tools': [s['function'] for s in schemas],
        'binding_identity':{'owner':owner,'workspace':context.runtime.workspace.workspace_id,'conversation':context.session_id,
            'reading_material':context.metadata.get('reading_material_id'),'mastery_path':context.metadata.get('mastery_path_id')},
        'system_prompt': system, 'history': json.dumps(history, ensure_ascii=False),
        'context': json.dumps({'messages': current_messages, 'source_references': context.metadata.get('source_references', [])}, ensure_ascii=False),
        'images': images,
        'language': context.language or 'en', 'output_dir': str(context.runtime.workspace.output_dir)}
    started = time.monotonic()
    events = []
    visible = {}
    # A finish guard is not a request to hide normal teaching. Mastery streams
    # its explanation before posting a question; only private protocol surfaces
    # explicitly declare buffers_visible_output.
    buffer_answer = pipeline._capability_buffers_visible_output(context)

    async def emit_text(stream_id, text):
        if buffer_answer:
            return
        previous = visible.get(stream_id, '')
        if text.startswith(previous) and len(text) > len(previous):
            await stream.content(text[len(previous):], source=pipeline.event_source, stage='responding',
                metadata={'call_id': stream_id, 'call_kind': 'agent_loop_round',
                    'trace_role': 'response', 'trace_kind': 'llm_chunk', 'execution_engine': 'harness'})
            visible[stream_id] = text

    async def on_live(response):
        await emit_text(response['id'], response.get('text', ''))

    async def on_question(record):
        from deeptutor.agents.loop.pipeline import _normalise_user_reply
        replies = []
        questions = record.get('questions', [])
        for start in range(0, len(questions), 4):
            batch = questions[start:start+4]
            captured = []
            waiter = context.runtime.wait_for_user_reply
            if not callable(waiter):
                raise RuntimeError('USER_INPUT_UNAVAILABLE: Reopen this conversation to answer the pending question.')
            async def capture_reply():
                reply = await waiter()
                captured.append(reply)
                return reply
            context.runtime.wait_for_user_reply = capture_reply
            try:
                await _invoke(entry, 'ask_user', {'questions':[
                    {'id':q['id'], 'prompt':q.get('question') or q.get('prompt'),
                     'header':q.get('header'), 'options':q.get('options', []),
                     'multi_select':bool(q.get('multiSelect')), 'allow_free_text':True} for q in batch]},
                    record['id'] + ':' + str(start))
            finally:
                context.runtime.wait_for_user_reply = waiter
            if not captured or captured[0] is None:
                return {'cancel': True}
            text, answers = _normalise_user_reply(captured[0])
            by_id = {a['questionId']:a['text'] for a in answers or []}
            for q in batch:
                answer = by_id.get(q['id'], text if len(batch) == 1 else '')
                exact = any(o.get('label') == answer for o in q.get('options', []))
                replies.append({'id':q['id'], 'selected':[answer] if exact else [],
                    **({} if exact else {'custom':answer})})
        return {'answers':replies}

    async def on_event(event):
        events.append(event)
        typ = event.get('type', '')
        # Domain operations already emit native tool cards; core operations use
        # the same event protocol so each entry retains live execution progress.
        if typ == 'tool/call' and event.get('tool_name') != 'invoke_surface_tool':
            await stream.tool_call(event.get('tool_name', 'Harness'), event.get('arguments') or {}, source=pipeline.event_source, stage='executing')
        elif typ == 'tool/result' and event.get('tool_name') != 'invoke_surface_tool':
            await stream.emit(StreamEvent(type=StreamEventType.TOOL_RESULT, source=pipeline.event_source,
                stage='executing', content=str(event.get('text') or event.get('result') or ''), metadata={'harness_event': event}))
        elif typ == 'assistant/message':
            await emit_text(event.get('stream_id') or str(event.get('seq')), event.get('text', ''))
        elif typ == 'deliverables/presented':
            from deeptutor.services.workspace import get_content_workspace_service
            service = get_content_workspace_service()
            binding = service.binding_by_id(context.runtime.workspace.workspace_id)
            items = []
            for item in event.get('files', []):
                # Host validates a private immutable snapshot. Copy that exact
                # revision into this turn's authorized delivery folder, including
                # when the user requested an output outside the workspace.
                path = Path(item.get('snapshot_path') or '').resolve()
                if not item.get('delivery_id') or not item.get('snapshot_path'):
                    raise ValueError('DELIVERY_SNAPSHOT_REQUIRED: Present this file again to create a verified attachment.')
                metadata = await asyncio.to_thread(validate_file, path)
                if metadata['sha256'] != item.get('sha256') or metadata['size_bytes'] != item.get('size_bytes'):
                    raise ValueError('DELIVERY_SNAPSHOT_CHANGED: The verified output bytes do not match.')
                # A compact workspace-local staging location avoids stacking a
                # full turn path, hash and UUID temporary suffix on Windows.
                # The authored output stays in the user's requested directory.
                target = Path(binding.root) / 'outputs' / '.harness-deliveries' / metadata['sha256'] / Path(item['filename']).name
                if not target.resolve().is_relative_to(Path(binding.root).resolve()):
                    raise ValueError('DELIVERY_SCOPE_CHANGED: The current output directory is unavailable.')
                def stage_snapshot():
                    from scripts.repair_bundle import atomic
                    data = path.read_bytes()
                    if hashlib.sha256(data).hexdigest() != metadata['sha256']:
                        raise ValueError('DELIVERY_SNAPSHOT_CHANGED')
                    atomic(target, data)
                await asyncio.to_thread(stage_snapshot)
                items.append({'path': target.relative_to(binding.root).as_posix()})
            if items:
                # Dispatch owns native tool_result metadata and attachment cards.
                # Calling the low-level executor alone silently drops that UI
                # contract even though the Host has validated the actual file.
                await _invoke(entry, 'workspace_present', {'items':items}, 'harness-delivery-' + str(event['seq']))

    task = None
    cancelled = False
    try:
        selection = selection_for_config(pipeline.llm_config)
        task = None
        for correction in range(2):
            text = (current_messages[-1].get('content') or context.user_message) if correction == 0 else pipeline._capability_finish_instruction(context, task.get('result', ''))
            if not text:
                break
            task = await ExecutionClient().execute(text=text, workspace_id=context.runtime.workspace.workspace_id,
                source=surface, task_key=stable_task_key(context.session_id, context.runtime.turn_id, text),
                llm_selection=selection, owner=owner, on_event=on_event, conversation_id=context.session_id,
                surface=surface, interaction=interaction, on_live=on_live, on_question=on_question,
                on_task=lambda value: entry.update(harness_task=value),
                source_references=context.metadata.get('source_references', []),
                expected_model=resolved_model_snapshot(pipeline.llm_config))
            if task['status'] != 'completed':
                error = task.get('error') or {'code': task.get('error_code') or task['status'], 'message': task.get('recovery') or 'Execution stopped.'}
                await stream.emit(StreamEvent(type=StreamEventType.ERROR, source=pipeline.event_source,
                    content=error.get('message', 'Execution failed'), metadata={'execution_error': error, 'task_id': task['id'],
                        'turn_terminal': True, 'status': 'cancelled' if task['status'] == 'cancelled' else 'failed',
                        'error_code': error.get('code', 'HARNESS_FAILED'), 'retryable': bool(error.get('retryable', False))}))
                return LoopOutcome(final_text=error.get('message', ''), completed=False)
            if 'terminal_override' in entry or not pipeline._capability_finish_instruction(context, task.get('result', '')):
                break
        final = task.get('result', '') if task else ''
        if 'terminal_override' not in entry and pipeline._capability_finish_instruction(context, final):
            raise RuntimeError('CAPABILITY_FINISH_REJECTED: Required teaching or reading steps have not completed.')
        override = entry['terminal_override'] if 'terminal_override' in entry else pipeline._capability_final_text_override(context, final)
        if override is not None:
            final = override
        if buffer_answer or not visible:
            await stream.content(final, source=pipeline.event_source, stage='responding',
                metadata={'call_id': 'harness-final-' + task['id'], 'call_kind': 'llm_final_response',
                    'trace_role': 'response', 'trace_kind': 'llm_output', 'execution_engine': 'harness'})
        context.capability_output.answer_published = True
        await stream.progress('', source=pipeline.event_source, stage='responding', metadata={
            'harness_task_id': task['id'], 'harness_session_id': task.get('session_id'),
            'execution_engine': 'harness', 'total_seconds': round(time.monotonic() - started, 3)})
        return LoopOutcome(final_text=final, completed=True)
    except asyncio.CancelledError:
        cancelled = True
        raise
    finally:
        # Polling transport failure is not task completion. Retain its tool
        # binding until an authenticated reconnect replaces it or the Host has
        # a confirmed terminal status; otherwise a live tool would see 409.
        if cancelled or task is not None and task.get('status') in {'completed', 'failed', 'cancelled', 'interrupted'}:
            _contexts.pop(ref, None)
            for operation in entry['calls'].values():
                if not operation.done():
                    operation.cancel()
        else:
            entry['cleanup'] = asyncio.create_task(_release_context_when_settled(ref, entry))
